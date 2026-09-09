"""Validazione manuale e riproducibile della pipeline JET sul caso Nordson.

Non è raccolto da pytest e non scrive dati o risultati del cliente su disco.
"""

from __future__ import annotations

import argparse
from datetime import time
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from backend.jet import (
    ParametriClienteJet,
    calcola_frequenza_conti,
    leggi_righe_xlsx,
    mappa_righe_giornale,
    valuta_riga,
    verifica_sequenza,
)


def _liste_cliente(percorso: Path) -> tuple[list[str], list[str], None]:
    workbook = load_workbook(percorso, read_only=True, data_only=True)
    try:
        foglio = workbook["Data Input"]
        staff = [
            str(foglio[f"O{riga}"].value).strip()
            for riga in range(3, 41)
            if foglio[f"O{riga}"].value not in (None, "")
        ]
        # Calcs1!W:AF referenzia esattamente L18:L27; L28 non è incluso
        # nelle formule del criterio, anche se contiene un ulteriore valore.
        parole = [
            str(foglio[f"L{riga}"].value).strip()
            for riga in range(18, 28)
            if foglio[f"L{riga}"].value not in (None, "")
        ]
        # Le formule del lookup puntano a L36:M53, ma L36:L53 contiene
        # formule =J36...=J53 e le celle sorgente J36:J53 sono vuote.
        festivita = None
        return staff, parole, festivita
    finally:
        workbook.close()


def _conti_original_data(percorso: Path) -> tuple[list[str | None], int]:
    """Legge il codice conto verificato sull'intera popolazione Nordson."""
    workbook = load_workbook(percorso, read_only=True, data_only=True)
    try:
        foglio = workbook["Original data"]
        intestazioni = next(foglio.iter_rows(values_only=True))
        indice = intestazioni.index("Conto n.")
        conti = [
            None
            if riga[indice] in (None, "")
            else (
                str(int(riga[indice]))
                if isinstance(riga[indice], float) and riga[indice].is_integer()
                else str(riga[indice]).strip()
            )
            for riga in foglio.iter_rows(min_row=2, values_only=True)
            # Le ultime colonne contengono formule trascinate anche oltre i
            # dati sorgente: una riga esiste solo se A:N contiene un valore.
            if any(valore not in (None, "") for valore in riga[:14])
        ]
        return conti, indice
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()

    staff, parole, festivita = _liste_cliente(args.workbook)
    grezze = leggi_righe_xlsx(args.workbook, foglio="Data Input")
    righe = mappa_righe_giornale(
        grezze,
        {
            "data_effettiva": "EffectiveDate",
            "data_creazione": "CreatedDate",
            "ora_creazione": "CreatedTime",
            "identificativo_registrazione": "TransactionId",
            # Data Input non conserva Doc.No.: numero_documento resta None.
            "importo_netto": "Net",
            "descrizione": "JournalDescription",
            "utente": "UserId",
        },
    )
    conti, indice_conto = _conti_original_data(args.workbook)
    if len(conti) != len(righe):
        raise ValueError(
            "Original data e Data Input non hanno lo stesso numero di righe: "
            "non è sicuro associare i conti per posizione"
        )
    righe = [
        riga.model_copy(update={"conto_contabile": conto})
        for riga, conto in zip(righe, conti)
    ]
    frequenze_conti = calcola_frequenza_conti(righe)
    parametri = ParametriClienteJet(
        materialita_bilancio=Decimal("1050000"),
        performance_materiality=Decimal("735000"),
        utile_netto_dopo_imposte=Decimal("9582473"),
        valore_medio_registrazione=Decimal("17207.03"),
        orario_ufficio_inizio=time(8, 30),
        orario_ufficio_fine=time(17),
        giorni_weekend=[5, 6],
        soglia_backdating_giorni=60,
        festivita=festivita,
        staff_autorizzato=staff,
        utenti_di_sistema=["BATCHJOB", "BANKBATCH"],
        parole_chiave_parti_correlate=parole,
        soglia_frequenza_insolita=None,
        conti_infragruppo_parte_correlata=None,
        soglia_da_investigare=4,
        punteggio_profit_impact=1,
        punteggio_oltre_dieci_volte_media=1,
        punteggio_sopra_performance_materiality=1,
        punteggio_importo_cifra_tonda=1,
        punteggio_weekend=1,
        punteggio_festivita=4,
        punteggio_fuori_orario=1,
        punteggio_backdated=4,
        punteggio_staff_non_autorizzato=4,
        punteggio_parte_correlata=4,
        punteggio_descrizione_vuota=4,
        punteggio_conto_insolito_raro=None,
        punteggio_conto_infragruppo_parte_correlata=None,
    )
    esiti = [valuta_riga(riga, parametri, frequenze_conti) for riga in righe]
    sequenza = verifica_sequenza(righe)

    print(f"righe={len(righe)}")
    print(f"staff={len(staff)} parole_chiave={len(parole)} festivita=non_disponibili")
    print(f"da_investigare={sum(esito.da_investigare for esito in esiti)}")
    senza_staff = sum(
        esito.punteggio_totale
        - (parametri.punteggio_staff_non_autorizzato if esito.flag_staff_non_autorizzato else 0)
        >= parametri.soglia_da_investigare
        for esito in esiti
    )
    print(f"da_investigare_senza_punti_staff={senza_staff}")
    print(
        "utenti_di_sistema_esclusi="
        f"{sum(riga.utente in {'BATCHJOB', 'BANKBATCH'} for riga in righe)}"
    )
    print(f"staff_non_autorizzato={sum(esito.flag_staff_non_autorizzato is True for esito in esiti)}")
    print(f"sopra_pm={sum(esito.flag_sopra_performance_materiality for esito in esiti)}")
    print(f"cifra_tonda={sum(esito.flag_importo_cifra_tonda for esito in esiti)}")
    print(f"buchi_sequenza={len(sequenza)}")
    print(f"conti_distinti={len(frequenze_conti)}")
    print(
        "conti_frequenza_1="
        f"{sum(frequenza == 1 for frequenza in frequenze_conti.values())}"
    )
    print(
        "conti_frequenza_sotto_10="
        f"{sum(frequenza < 10 for frequenza in frequenze_conti.values())}"
    )
    print(
        "conti_frequenza_sotto_5="
        f"{sum(frequenza < 5 for frequenza in frequenze_conti.values())}"
    )
    print(
        "righe_su_conti_frequenza_sotto_5="
        f"{sum(frequenza for frequenza in frequenze_conti.values() if frequenza < 5)}"
    )
    print(
        "conti_frequenza_sopra_1000="
        f"{sum(frequenza > 1000 for frequenza in frequenze_conti.values())}"
    )
    print(f"conto_frequenza_massima={max(frequenze_conti.values(), default=0)}")
    totale_conti = len(conti)
    conti_vuoti = sum(conto is None for conto in conti)
    conti_numerici = sum(conto is not None and conto.isdigit() for conto in conti)
    conti_testuali = totale_conti - conti_vuoti - conti_numerici
    print(f"diagnostica_conto_colonna=Conto n. indice={indice_conto + 1}")
    print(f"diagnostica_conto_numerici_pct={conti_numerici / totale_conti:.4%}")
    print(f"diagnostica_conto_testuali_pct={conti_testuali / totale_conti:.4%}")
    print(f"diagnostica_conto_vuoti_pct={conti_vuoti / totale_conti:.4%}")
    print(f"diagnostica_conto_valori_distinti={len(frequenze_conti)}")


if __name__ == "__main__":
    main()
