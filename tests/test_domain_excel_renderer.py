"""Fase 4 — renderer Excel come serializzatore dei risultati di dominio."""

from pathlib import Path

from openpyxl import load_workbook

from backend.catalog import MASTER_XLSX, checklist_items
from backend.domain import (
    Evidence,
    ExtractedField,
    PraticaRecord,
    VerificationResult,
    documents_to_evidence,
    evaluate_pratica,
)
from backend.domain.client_config_loader import load_client_config
from backend.domain.client_classify import scan_folder_for_client
from backend.domain.excel_renderer import INDEX_CELLS, render_domain_workbook

ROOT = Path(__file__).resolve().parents[1]
FIXTURES_DOCS = ROOT / "fixtures" / "docs"


def test_renderer_writes_domain_results_and_preserves_datasnipper_sheets(tmp_path: Path):
    pratica = PraticaRecord(
        id="render-test",
        client_id="demo",
        client="Cliente Demo",
        period="Aprile - Giugno 2026",
        documents_dir=str(FIXTURES_DOCS),
    )
    config = load_client_config("demo")
    documents = scan_folder_for_client(str(FIXTURES_DOCS), pratica.id, config)
    evidences = documents_to_evidence(pratica.id, documents, config)
    verifiche = evaluate_pratica(
        pratica_id=pratica.id,
        client=pratica.client,
        period=pratica.period,
        evidences=evidences,
        client_config=config,
    )

    template = load_workbook(MASTER_XLSX)
    ds_rows = {
        name: template[name].max_row
        for name in template.sheetnames
        if name.startswith("DS_INTERNAL_")
    }
    output = render_domain_workbook(pratica, verifiche, [], tmp_path)
    workbook = load_workbook(output)

    assert output.is_file()
    assert [workbook["INDICE"][INDEX_CELLS[s]].value for s in INDEX_CELLS] == [
        verifiche[s].status for s in INDEX_CELLS
    ]
    assert {
        name: workbook[name].max_row for name in ds_rows
    } == ds_rows

    rows = {item["id"]: item["row"] for item in checklist_items()}
    assert workbook["Richiesta doc"][f"Q{rows['B.1']}"].value == "✓"
    assert workbook["Richiesta doc"][f"Q{rows['A.1']}"].value == "wip"
    assert "DOMINIO" in workbook.sheetnames
    assert "**A.1**" in (tmp_path / "mancanti.md").read_text(encoding="utf-8")


def test_renderer_includes_extracted_fields_in_domain_detail(tmp_path: Path):
    pratica = PraticaRecord(
        id="fields-test",
        client_id="demo",
        client="Cliente Demo",
        period="Aprile - Giugno 2026",
        documents_dir="/documenti",
    )
    evidence = Evidence(
        pratica_id=pratica.id,
        item_id="E.1",
        source_name="F24_aprile_2026.txt",
        source_path="/documenti/F24_aprile_2026.txt",
        method="txt",
        fields=[
            ExtractedField(kind="importo", value="35616.68", unit="EUR"),
            ExtractedField(kind="data_versamento", value="16/04/2026"),
        ],
    )
    verifiche = {
        "C": VerificationResult(
            pratica_id=pratica.id,
            client=pratica.client,
            period=pratica.period,
            section="C",
            status="wip",
            reasoning="Evidenza parziale.",
            evidence=[evidence],
        )
    }

    output = render_domain_workbook(pratica, verifiche, [], tmp_path)
    sheet = load_workbook(output)["DOMINIO"]
    evidence_row = next(row for row in sheet.iter_rows(min_row=2) if row[2].value == "Evidenza")

    assert evidence_row[5].value == (
        "/documenti/F24_aprile_2026.txt\n"
        "importo: 35616.68 EUR · data versamento: 16/04/2026"
    )
