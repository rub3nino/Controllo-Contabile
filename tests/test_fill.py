from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from backend.classify import scan_folder
from backend.fill import Filler
from backend.models import DocumentOut, PraticaIn
from backend.provenance import ProvenanceLog

DS_SHEETS = (
    "DS_INTERNAL_DOCUMENT_STORAGE",
    "DS_INTERNAL_SETTINGS_STORAGE",
    "DS_INTERNAL_DOCGROUP_STORAGE",
    "DS_INTERNAL_SNIP_STORAGE",
)


def _pratica(docs_dir: Path) -> PraticaIn:
    return PraticaIn(
        client="Cliente Demo",
        period="Aprile - Giugno 2026",
        done_by="RH",
        documents_dir=str(docs_dir),
        activity_date="2026-09-08",
    )


def test_scan_and_fill(tmp_path: Path | None = None):
    docs_dir = ROOT / "fixtures" / "docs"
    docs = scan_folder(str(docs_dir))
    by_name = {d.name: d.item_id for d in docs}
    assert by_name["F24_aprile_2026.txt"] == "E.1", by_name
    assert by_name["Estratto_Intesa_30.06.2026.txt"] == "F.1", by_name
    assert by_name["Bilancino_30.06.2026.csv"] == "B.1", by_name
    assert by_name["Verbale_CDA_05.05.2026.txt"] == "C.2", by_name
    assert by_name["LIPE_2trim_2026.txt"] == "E.3", by_name
    assert by_name["Registro_IVA_acquisti.txt"] == "D.3", by_name

    out = tmp_path or (ROOT / "output" / "_test")
    pratica = _pratica(docs_dir)
    prov = ProvenanceLog(out / "provenienza.jsonl")
    filler = Filler(pratica, docs, out, prov)
    xlsx = filler.fill_all([], ["E.5"])
    assert xlsx.exists()
    assert filler.checklist["E.1"] == "✓"
    assert filler.checklist["F.1"] == "✓"
    assert filler.checklist["B.1"] == "✓"
    assert filler.checklist["A.1"] == "wip"
    assert filler.checklist["E.5"] == "N/A"
    assert any(r.sheet == "C" and r.cell.endswith("12") for r in prov.rows)
    assert any(r.sheet == "E" and r.cell.startswith("F") for r in prov.rows)
    assert any("check bancario" in w for w in filler.warnings)
    assert (out / "mancanti.md").exists()
    missing_ids = {m["id"] for m in filler.missing}
    assert "E.5" not in missing_ids
    assert "E.1" not in missing_ids

    wb = load_workbook(xlsx)
    for name in DS_SHEETS:
        assert name in wb.sheetnames
    master = load_workbook(ROOT / "Template_MASTER.xlsx")
    assert wb[DS_SHEETS[0]].max_row == master[DS_SHEETS[0]].max_row
    assert wb["INDICE"]["F10"].value == "wip"  # A: giudizio
    assert wb["INDICE"]["F13"].value == "wip"  # D: giudizio
    assert wb["INDICE"]["F17"].value == "wip"  # H
    assert wb["INDICE"]["F18"].value == "wip"  # I
    assert wb["INDICE"]["F14"].value == "wip"  # E: manca F.2/F.3
    assert wb["Richiesta doc"]["Q29"].value == "✓"  # E.1 row 29
    qfill = wb["Richiesta doc"]["Q29"].fill
    assert qfill.fill_type == "solid"
    c = wb["C"]
    assert c["B10"].value.month == 4
    assert c["B11"].value.day == 16 and c["B11"].value.month == 4
    assert abs(float(c["B12"].value) - 35616.68) < 0.01
    assert c["C24"].value == "✓"  # LIPE on 2° trim
    assert "PROVENIENZA" in wb.sheetnames
    print("ok", xlsx, "rows", len(prov.rows), "missing", len(filler.missing), "warn", filler.warnings)


def test_judgment_never_auto_done(tmp_path: Path):
    tmp_path.mkdir(parents=True, exist_ok=True)
    docs_dir = ROOT / "fixtures" / "docs"
    docs = scan_folder(str(docs_dir))
    docs.append(
        DocumentOut(
            id="fake-a1",
            name="SCI.txt",
            path=str(docs_dir / "SCI.txt"),
            ext=".txt",
            size=1,
            item_id="A.1",
            item_label="SCI",
            confidence=1,
            method="umano",
        )
    )
    docs.append(
        DocumentOut(
            id="fake-a2",
            name="ops.txt",
            path=str(docs_dir / "ops.txt"),
            ext=".txt",
            size=1,
            item_id="A.2",
            item_label="ops",
            confidence=1,
            method="umano",
        )
    )
    docs.append(
        DocumentOut(
            id="fake-a3",
            name="cartelle.txt",
            path=str(docs_dir / "cartelle.txt"),
            ext=".txt",
            size=1,
            item_id="A.3",
            item_label="cartelle",
            confidence=1,
            method="umano",
        )
    )
    pratica = _pratica(docs_dir)
    pratica.na_items = ["A.1", "A.2", "A.3", "A.4"]
    prov = ProvenanceLog(tmp_path / "provenienza.jsonl")
    filler = Filler(pratica, docs, tmp_path, prov)
    xlsx = filler.fill_all([], pratica.na_items)
    wb = load_workbook(xlsx)
    assert wb["INDICE"]["F10"].value == "wip"
    assert wb["INDICE"]["F18"].value == "wip"
    assert "Evidenza in cartella" in str(wb["A"]["A9"].value)
    assert wb["INDICE"]["F10"].value != "✗"


def test_skip_section_writes_cross(tmp_path: Path):
    tmp_path.mkdir(parents=True, exist_ok=True)
    docs_dir = ROOT / "fixtures" / "docs"
    docs = scan_folder(str(docs_dir))
    pratica = _pratica(docs_dir)
    pratica.skip_sections = ["A", "I"]
    prov = ProvenanceLog(tmp_path / "p2.jsonl")
    filler = Filler(pratica, docs, tmp_path, prov)
    xlsx = filler.fill_all([], [], pratica.skip_sections)
    wb = load_workbook(xlsx)
    assert wb["INDICE"]["F10"].value == "✗"
    assert wb["INDICE"]["F18"].value == "✗"
    assert wb["INDICE"]["F11"].value != "✗"


def test_fondi_go_to_named_quarter(tmp_path: Path):
    tmp_path.mkdir(parents=True, exist_ok=True)
    docs_dir = ROOT / "fixtures" / "docs"
    docs = scan_folder(str(docs_dir))
    for name, iid in (("ENASARCO_1TR26.pdf", "E.2"), ("PREVIMODA_2TR26.pdf", "E.2")):
        docs.append(
            DocumentOut(
                id=name,
                name=name,
                path=str(docs_dir / "LIPE_2trim_2026.txt"),
                ext=".pdf",
                size=1,
                item_id=iid,
                item_label="fondi",
                confidence=1,
                method="filename",
            )
        )
    pratica = _pratica(docs_dir)
    pratica.period = "II trimestre 2026"
    filler = Filler(pratica, docs, tmp_path, ProvenanceLog(tmp_path / "p3.jsonl"))
    xlsx = filler.fill_all([], [])
    wb = load_workbook(xlsx)
    ws = wb["C"]
    by_name = {str(ws[f"A{r}"].value): r for r in range(19, 24)}
    assert "ENASARCO_1TR26" in by_name
    assert "PREVIMODA_2TR26" in by_name
    assert ws[f"B{by_name['ENASARCO_1TR26']}"].value == "✓"
    assert ws[f"C{by_name['PREVIMODA_2TR26']}"].value == "✓"
    assert ws["C24"].value == "✓"


def test_giornale_last_row(tmp_path: Path):
    tmp_path.mkdir(parents=True, exist_ok=True)
    from openpyxl import Workbook

    from backend.extract import parse_giornale_last

    path = tmp_path / "giornale.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Gaglianico Pagina 12"])
    ws.append([None, "Numero", "Data acq.", "N. doc.", None, None, "Periodo", None, "Data reg."])
    ws.append([None, 21698, "07.07.2026", 7770006914, None, None, 6, None, "30.06.2026", None, "30.06.2026", "AR", "Fattura clienti"])
    wb.save(path)
    meta = parse_giornale_last(path)
    assert meta["nr_reg"] == 21698
    assert "30.06.2026" in str(meta["data_reg"])
    assert meta["descrizione"] == "Fattura clienti"
    assert meta["page"] == 12


if __name__ == "__main__":
    from tempfile import TemporaryDirectory

    from backend.extract import find_f24_importo, find_payment_date

    ocr_sample = "QUIETANZA\n629.91147\n18052026\nDATADELVERSAMENTO\nERARIO\n629.911,47\n"
    assert find_f24_importo(ocr_sample) == 629911.47
    assert find_payment_date(ocr_sample) == "18/05/2026"

    test_scan_and_fill()
    with TemporaryDirectory() as d:
        root = Path(d)
        test_judgment_never_auto_done(root / "j")
        test_skip_section_writes_cross(root / "skip")
        test_fondi_go_to_named_quarter(root / "fondi")
        test_giornale_last_row(root / "gio")
    print("all ok")
