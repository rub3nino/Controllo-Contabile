"""Build Template_MASTER.xlsx from the Ferrero instance. Original file is not overwritten."""
from openpyxl import load_workbook

SRC = "/Users/ruben/Progetti/Controllo Contabile automatizzato/Template GF wps verifica trimestrale II TRIM .xlsx"
DST = "/Users/ruben/Progetti/Controllo Contabile automatizzato/Template_MASTER.xlsx"

DATE_FMT = "mm-dd-yy"
RICHIESTA_DATA_ROWS = [5, 6, 7, 8, 11, 12, 13, 14, 17, 18, 19, 20, 24, 25, 26, 29, 30, 31, 32, 33, 36, 37, 38, 41, 42]
RICHIESTA_CLEAR_COLS = list("DEFGIJKLNOPQ")


def clear_value(cell):
    from openpyxl.cell.cell import MergedCell
    if isinstance(cell, MergedCell):
        return
    cell.value = None


def set_header_formulas(ws, include_reviewed=True, include_date=True):
    ws["B1"].value = "=+INDICE!A1"
    ws["B2"].value = "=+INDICE!M1"
    ws["B3"].value = "=INDICE!C28"
    if include_reviewed and ws["A4"].value:
        ws["B4"].value = "=+INDICE!C29"
    if include_date and ws["A5"].value:
        ws["B5"].value = "=INDICE!F6"
        ws["B5"].number_format = DATE_FMT


def drop_images(ws):
    if getattr(ws, "_images", None):
        ws._images = []


def drop_ole(ws):
    for attr in ("_ole_objects", "ole_objects"):
        if hasattr(ws, attr):
            try:
                setattr(ws, attr, [])
            except Exception:
                pass


wb = load_workbook(SRC)

# --- INDICE ---
idx = wb["INDICE"]
idx["A1"] = "[CLIENTE]"
idx["M1"] = "[PERIODO]"
idx["F5"] = None
idx["F5"].number_format = DATE_FMT
idx["F6"] = None
idx["F6"].number_format = DATE_FMT
idx["C28"] = None
for row in range(10, 19):
    idx[f"F{row}"] = None

# --- Richiesta doc ---
rd = wb["Richiesta doc"]
rd["D3"] = "[T-2]"
rd["I3"] = "[T-1]"
rd["N3"] = "[Trimestre corrente]"
for row in RICHIESTA_DATA_ROWS:
    for col in RICHIESTA_CLEAR_COLS:
        clear_value(rd[f"{col}{row}"])

# --- A ---
a = wb["A"]
set_header_formulas(a)
drop_images(a)

# --- B ---
b = wb["B"]
set_header_formulas(b)
for row in range(11, 16):
    for col in "CDEFG":
        clear_value(b[f"{col}{row}"])

# --- C ---
c = wb["C"]
set_header_formulas(c)
for col in "BCDEFG":
    for row in (10, 11, 12, 13):
        clear_value(c[f"{col}{row}"])
        if row in (10, 11):
            c[f"{col}{row}"].number_format = DATE_FMT if row == 11 else "[$-410]mmm\\-yy;@"
for col in "BCDE":
    for row in range(19, 25):
        clear_value(c[f"{col}{row}"])
# fondi names except LIPE are client-specific
for row, label in [(19, "[FONDO]"), (20, "[FONDO]"), (21, "[FONDO]"), (22, "[FONDO]"), (23, "[FONDO]")]:
    c[f"A{row}"] = label
c["C28"] = None
c["D28"] = None
c["B27"] = None

# --- D ---
d = wb["D"]
set_header_formulas(d)
d["A8"] = None

# --- E ---
e = wb["E"]
set_header_formulas(e)
e["A9"] = "RICONCILIAZIONI BANCARIE"
for row in range(12, 29):
    for col in "ABCDEFGI":
        clear_value(e[f"{col}{row}"])
    e[f"H{row}"] = f"=E{row}-F{row}"
    e[f"H{row}"].number_format = '_-* #,##0.00\\ [$€-410]_-;\\-* #,##0.00\\ [$€-410]_-;_-* "-"??\\ [$€-410]_-;_-@_-'
    e[f"E{row}"].number_format = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
    e[f"F{row}"].number_format = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
    e[f"G{row}"].number_format = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'

# --- F ---
f = wb["F"]
set_header_formulas(f)
for row in range(10, 14):
    for col in "CDEF":
        clear_value(f[f"{col}{row}"])
drop_ole(f)

# --- G ---
g = wb["G"]
set_header_formulas(g)
g["E9"] = "Saldo periodo"
g["F9"] = "Saldo N-1"
g["F9"].number_format = DATE_FMT

# --- H / I ---
for name in ("H", "I"):
    set_header_formulas(wb[name])

# --- DataSnipper: strip Ferrero document payloads, keep sheets ---
for name in (
    "DS_INTERNAL_DOCUMENT_STORAGE",
    "DS_INTERNAL_SETTINGS_STORAGE",
    "DS_INTERNAL_DOCGROUP_STORAGE",
    "DS_INTERNAL_SNIP_STORAGE",
):
    ws = wb[name]
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.value = None

wb.properties.title = "Controllo contabile trimestrale — MASTER"
wb.properties.subject = "Template vuoto per tutti i clienti"

wb.save(DST)
print("wrote", DST)
