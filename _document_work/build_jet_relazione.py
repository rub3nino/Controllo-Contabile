from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path("deliverables")
OUT.mkdir(exist_ok=True)
DOCX = OUT / "Relazione_analisi_JET_AS-IS_TO-BE.docx"

BLUE = "235A78"
DARK = "183245"
LIGHT = "E8F0F5"
GRAY = "F2F4F7"
GREEN = "E6F3EA"
AMBER = "FFF4D6"
RED = "FBE8E7"
MUTED = RGBColor(92, 105, 115)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.82)
sec.bottom_margin = Inches(0.78)
sec.left_margin = Inches(0.86)
sec.right_margin = Inches(0.86)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.3)
normal.font.color.rgb = RGBColor(31, 43, 52)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

for name, size, color, before, after in [
    ("Title", 25, DARK, 0, 8),
    ("Subtitle", 12.5, "4E6574", 0, 14),
    ("Heading 1", 16, BLUE, 15, 7),
    ("Heading 2", 13, BLUE, 11, 5),
    ("Heading 3", 11.5, DARK, 8, 3),
]:
    s = styles[name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.font.bold = name != "Subtitle"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

for lname in ["List Bullet", "List Number"]:
    s = styles[lname]
    s.font.name = "Calibri"
    s.font.size = Pt(10.3)
    s.paragraph_format.left_indent = Inches(0.28)
    s.paragraph_format.first_line_indent = Inches(-0.18)
    s.paragraph_format.space_after = Pt(3)
    s.paragraph_format.line_spacing = 1.08

if "Control label" not in styles:
    s = styles.add_style("Control label", WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = "Calibri"
    s.font.size = Pt(9)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(BLUE)
    s.paragraph_format.space_before = Pt(3)
    s.paragraph_format.space_after = Pt(1)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=85, start=110, bottom=85, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths))); tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "0"); tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for x in list(grid): grid.remove(x)
    for w in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(w)); grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i]/1440)
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths[i])); tcW.set(qn("w:type"), "dxa")
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_cell_text(cell, text, bold=False, color=None, size=8.6, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.02
    if align is not None: p.alignment = align
    r = p.add_run(str(text))
    r.bold = bold; r.font.name = "Calibri"; r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if color: r.font.color.rgb = RGBColor.from_string(color)

def add_table(headers, rows, widths, font=8.6):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    set_table_geometry(t, widths)
    for i, h in enumerate(headers):
        shade(t.rows[0].cells[i], BLUE)
        set_cell_text(t.rows[0].cells[i], h, True, "FFFFFF", font, WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, False, None, font, WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.CENTER)
            if i == 1:
                fill = GREEN if str(val).startswith("Presente") and "parzial" not in str(val).lower() else AMBER if "Parziale" in str(val) or "diverso" in str(val).lower() else RED
                shade(cells[i], fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p

def add_label(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + ": ")
    r.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
    p.add_run(text)
    return p

def add_callout(title, text, fill=LIGHT):
    t = doc.add_table(rows=1, cols=1)
    set_table_geometry(t, [9760])
    t.style = "Table Grid"
    c = t.cell(0,0); shade(c, fill); margins(c, 125, 150, 125, 150)
    c.text = ""
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title); r.bold = True; r.font.color.rgb = RGBColor.from_string(DARK)
    p2 = c.add_paragraph(text); p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.line_spacing = 1.05
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def page_break():
    doc.add_page_break()

# Header/footer
hp = sec.header.paragraphs[0]
hp.text = "QUADRA  |  JET — RELAZIONE AS-IS / TO-BE"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = MUTED
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Documento per approvazione progettuale  •  10 settembre 2026")
r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = MUTED

# Cover / memo masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(5)
r = p.add_run("RELAZIONE TECNICO-FUNZIONALE")
r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph(style="Title")
p.add_run("Journal Entry Testing\nAS-IS e proposta TO-BE")
p = doc.add_paragraph(style="Subtitle")
p.add_run("Valutazione dei 15 controlli richiesti per l’evoluzione da MVP a prodotto operativo")

meta = [
    ("Destinatari", "Partner, manager e responsabili del progetto"),
    ("Perimetro", "Analisi funzionale e metodologica; nessuna modifica al software"),
    ("Base", "ISA Italia 240 §33 e §A44; riferimento tecnico JET; codice Quadra attuale"),
    ("Versione", "1.0 — 10 settembre 2026"),
]
for k,v in meta:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    rr=p.add_run(k+": "); rr.bold=True; rr.font.color.rgb=RGBColor.from_string(DARK)
    p.add_run(v)

add_callout("Messaggio decisionale", "L’MVP contiene una parte significativa dei controlli di preselezione, ma diversi requisiti richiesti sono assenti o implementati con una logica diversa. Prima dello sviluppo occorre approvare unità di analisi, soglie, punteggi, fonti dati e trattamento dei controlli non calcolabili.", AMBER)

doc.add_heading("Sintesi esecutiva", level=1)
doc.add_paragraph("Il JET attuale valuta prevalentemente singole righe del libro giornale. Ogni indicatore vero attribuisce un peso; la somma viene confrontata con la soglia “da investigare”. La selezione non certifica errore o frode: identifica registrazioni sulle quali svolgere follow-up documentale e professionale.")
add_bullet("Già presenti e sostanzialmente coerenti: impatto sull’utile, oltre 10× la media, superamento della performance materiality, fuori orario, descrizione vuota, conto raro, sequenza numerica.")
add_bullet("Presenti ma da correggere o ampliare: cifra tonda, weekend/festività, retrodatazione/fine periodo, utenti non autorizzati, keyword.")
add_bullet("Assenti: cifre finali ripetute, seconda analisi strutturata dei sospetti, lunghezza anomala del conto, calendari multinazionali automatici, OCR della visura e incrocio con poteri/accessi.")
add_bullet("Il test di sequenza e la successiva re-analisi non dovrebbero concorrere automaticamente al punteggio di rischio della singola riga.")

doc.add_heading("Decisioni richieste a partner e manager", level=2)
for x in [
    "Approvare se il JET debba selezionare righe, scritture complete o entrambe.",
    "Approvare una matrice dei pesi e la regola per gli indicatori obbligatori indipendenti dal punteggio.",
    "Stabilire chi è responsabile della validazione di calendario, utenti, parti correlate e data di chiusura.",
    "Stabilire come documentare criteri non eseguibili per dati mancanti.",
    "Separare l’analisi antifrode dai controlli di integrità e completezza della popolazione.",
]: add_bullet(x)

page_break()
doc.add_heading("1. Base normativa e metodo di valutazione", level=1)
doc.add_paragraph("L’ISA Italia 240 §33(a) richiede la verifica delle scritture contabili e delle rettifiche, con attenzione specifica alle registrazioni di fine periodo. Il §A44 fornisce caratteristiche di rischio — conti insoliti, autori inusuali, descrizioni scarse, importi tondi, operazioni infragruppo e fuori dal normale corso — ma non prescrive soglie numeriche o punteggi. Le formule quantitative derivano dal modello operativo NORDSON/BTI e dalle scelte progettuali di Quadra.")
add_callout("Principio di interpretazione", "Un indicatore vero significa “selezionare per approfondimento”, non “frode”. Un indicatore non calcolabile deve essere visibile come tale e non confuso con un risultato negativo.")

doc.add_heading("Scala utilizzata nella relazione", level=2)
rows = [
    ("Presente", "Il controllo esiste e la logica richiesta è sostanzialmente disponibile."),
    ("Parziale/diverso", "Esiste un controllo collegato, ma formula, perimetro o dati differiscono dal requisito."),
    ("Assente", "Non esiste oggi una logica che produca il risultato richiesto."),
]
add_table(["Stato", "Significato"], rows, [1900, 7860], 9)

doc.add_heading("Matrice generale dei 15 controlli", level=1)
matrix = [
    ("1", "Presente", "1", "Confermare formula e gestione utile nullo/perdita"),
    ("2", "Presente", "1", "Calcolare automaticamente la media assoluta"),
    ("3", "Presente", "1", "Importare/validare PM da Global Focus"),
    ("4", "Parziale/diverso", "1", "Oggi divisibile per 10; definire 10.000/100.000"),
    ("5", "Parziale/diverso", "1 + 1 oggi", "Calendari per Paese assenti; data usata da chiarire"),
    ("6", "Presente", "1", "Default 08:00–18:00 configurabile"),
    ("7", "Parziale/diverso", "1 oggi", "Manca finestra ultimi 5 giorni lavorativi"),
    ("8", "Parziale/diverso", "1 oggi", "Manca interruttore esplicito; dati spesso assenti"),
    ("9", "Presente", "1 oggi", "Nel modello NORDSON valeva 4"),
    ("10", "Presente", "opzionale/null", "Definire 2–4 utilizzi e base annuale"),
    ("11", "Presente separato", "0", "Solo sequenze numeriche; regole per pagina/serie"),
    ("12", "Assente", "da approvare", "Definire pattern e soglia importo"),
    ("13", "Assente", "0", "Workflow di secondo livello, non nuovo rischio"),
    ("14", "Assente", "da approvare", "Validare rispetto ai formati ERP"),
    ("15", "Parziale + assente OCR", "1 oggi keyword", "Separare keyword, visura, accessi e poteri"),
]
add_table(["#", "Stato attuale", "Peso attuale", "Intervento principale"], matrix, [500, 2100, 1400, 5760], 7.6)

page_break()
doc.add_heading("2. Analisi dettagliata dei controlli", level=1)

controls = [
{
"n":"1", "title":"Impatto superiore al 10% dell’utile netto", "status":"PRESENTE", "score":"Peso iniziale attuale: 1.",
"current":"Il sistema prende il valore assoluto della riga e lo confronta con il 10% del valore assoluto dell’utile netto dopo le imposte: |importo| > 10% × |utile|. L’uguaglianza non attiva il flag.",
"data":"Utile netto dopo imposte dello stesso periodo e importo netto correttamente normalizzato. Fonte: conto economico/bilancio di verifica e carta di lavoro di pianificazione.",
"expected":"Flag “Impatto sull’utile”, incremento del punteggio e motivazione visibile nell’export.",
"gap":"Con utile nullo o molto basso il test diventa eccessivamente sensibile; con perdita usa il valore assoluto. Occorre prevedere un benchmark alternativo o una motivazione del revisore.",
"proposal":"Mantenere il peso 1 come indicatore quantitativo, ma bloccare o qualificare il controllo quando l’utile non è un benchmark appropriato. Registrare fonte e periodo del valore inserito."},
{
"n":"2", "title":"Importi superiori a 10 volte la media", "status":"PRESENTE, CON MIGLIORIA NECESSARIA", "score":"Peso iniziale attuale: 1.",
"current":"Confronta |importo della riga| con 10 × |valore medio registrazione|. Il valore medio è oggi inserito manualmente.",
"data":"Popolazione completa del giornale, dopo pulizia e gestione duplicati. La media consigliata è somma degli importi assoluti divisa per numero di righe; non la media algebrica, che può annullarsi tra Dare e Avere.",
"expected":"Flag “>10× media” sulle righe eccezionalmente grandi rispetto alla popolazione.",
"gap":"La media manuale non è riproducibile; inoltre una media per riga può non rappresentare il valore della scrittura completa e risente fortemente degli outlier.",
"proposal":"Calcolo automatico e tracciato; mostrare numeratore, denominatore e metodo. Valutare anche media/mediana per tipo documento o conto, senza sostituire la regola approvata."},
{
"n":"3", "title":"Superamento della materialità da Global Focus", "status":"PRESENTE PER LA PERFORMANCE MATERIALITY", "score":"Peso iniziale attuale: 1.",
"current":"Il sistema confronta |importo| con la performance materiality (PM). La casella “materialità di bilancio” è salvata ma non alimenta alcun controllo.",
"data":"Performance materiality approvata per società, esercizio e perimetro. Fonte prevista: file Materiality/Global Focus, con verifica del valore, valuta, versione e data di approvazione.",
"expected":"Flag “Oltre performance materiality” e incremento del punteggio.",
"gap":"Non esiste importazione diretta o riconciliazione con Global Focus. “Materialità complessiva” e “PM” non devono essere confuse. La PM potrebbe variare per area o componente.",
"proposal":"Prevedere caricamento del file o inserimento manuale con evidenza della fonte; richiedere conferma esplicita di valuta e periodo. Mantenere peso 1 salvo diversa metodologia approvata."},
{
"n":"4", "title":"Cifre tonde divisibili per 10.000 o 100.000", "status":"PRESENTE MA CON FORMULA DIVERSA", "score":"Peso iniziale attuale: 1.",
"current":"Oggi una riga è “tonda” se |importo| è divisibile per 10. Quindi anche 20 o 150 vengono segnalati.",
"data":"Importo netto; eventuale soglia minima; valuta della registrazione. Se vi sono più valute, il multiplo deve essere interpretato nella valuta della riga o dopo conversione definita.",
"expected":"Distinguere almeno “multiplo di 10.000” e “multiplo di 100.000”, mostrando quale regola è scattata.",
"gap":"Il requisito proposto non coincide con il codice. Non è chiaro se 100.000 debba attivare entrambi i flag o solo quello più significativo, né come gestire decimali/valute.",
"proposal":"Parametri configurabili 10.000 e 100.000; un solo indicatore con livello massimo raggiunto per evitare doppio conteggio. Peso suggerito da sottoporre ad approvazione: 1, con soglia minima e motivazione."},
{
"n":"5", "title":"Weekend e festività per Paese", "status":"PARZIALE", "score":"Oggi due criteri distinti, peso iniziale 1 ciascuno.",
"current":"Weekend: verifica se la data effettiva cade nei giorni configurati. Festività: verifica se la data effettiva è nell’elenco inserito manualmente.",
"data":"Paese/sede applicabile alla società, calendario annuale, festività nazionali e locali, chiusure aziendali. Per Italia, Germania, Francia e Spagna occorre gestire anche festività regionali/locali.",
"expected":"Flag separati “weekend” e “festività”, con Paese/calendario e festività che ha causato la selezione.",
"gap":"Non esiste libreria/calendario automatico. Il codice usa la data effettiva: per sapere quando l’utente ha operato sarebbe normalmente più pertinente la data di creazione.",
"proposal":"Calendari versionati interni o libreria affidabile, con override cliente. Usare data creazione per attività fuori giorno lavorativo e, se utile, mantenere un indicatore distinto sulla data effettiva. Peso 1 ciascuno; evitare di contare due volte sabato che coincida con festività, salvo scelta esplicita."},
{
"n":"6", "title":"Registrazioni prima delle 08:00 o dopo le 18:00", "status":"PRESENTE", "score":"Peso iniziale attuale: 1.",
"current":"Confronta l’ora di creazione con inizio e fine configurati. Gli estremi sono inclusi: 08:00 e 18:00 non scattano; 07:59 e 18:01 scattano.",
"data":"Ora tecnica di creazione, fuso orario, orario ordinario della società e indicazione di account batch/turnisti.",
"expected":"Flag “Fuori orario” con ora rilevata e fascia applicata.",
"gap":"Non esiste oggi un default obbligatorio 08:00–18:00; è inserito dall’utente. Mancano fuso orario, turni, periodi di closing e orari diversi per sede.",
"proposal":"Proporre 08:00–18:00 come default modificabile, mai come regola normativa. Escludere o classificare separatamente batch; conservare sempre la fascia utilizzata nel risultato."},
{
"n":"7", "title":"Retrodatazione e ultimi 5 giorni lavorativi di chiusura", "status":"PARZIALE; CONTROLLO DI FINE PERIODO ASSENTE", "score":"Peso iniziale attuale: 1. Nel modello NORDSON il backdating valeva 4.",
"current":"Retrodatazione: data creazione − data effettiva ≥ soglia configurata. Non esiste un controllo sugli ultimi cinque giorni lavorativi prima della chiusura.",
"data":"Data effettiva, data/ora creazione, data di chiusura (default 31/12 modificabile), calendario lavorativo della società, data eventuale riapertura e tipo di scrittura.",
"expected":"Due esiti distinti: A) “retrodatata oltre N giorni”; B) “scrittura nella finestra di closing”, calcolata sugli ultimi 5 giorni lavorativi rispetto alla data di chiusura configurata. Opportuno includere anche scritture create dopo la chiusura con competenza precedente.",
"gap":"La richiesta combina due rischi diversi. La finestra di chiusura è un sottoinsieme da esaminare per disposizione ISA e non dovrebbe dipendere solo da un punteggio aggregato.",
"proposal":"Separare i controlli. Retrodatazione: peso forte da approvare (coerente con NORDSON: 4). Finestra closing: inclusione obbligatoria nel perimetro o lista dedicata, senza affidarsi al solo punteggio. Data default 31/12, sempre modificabile e documentata."},
{
"n":"8", "title":"Utenti non autorizzati con possibilità di disattivazione", "status":"PARZIALE", "score":"Peso iniziale attuale: 1; nel modello NORDSON valeva 4.",
"current":"Se esiste una whitelist, segnala l’utente non presente. Gli utenti di sistema sono esclusi dal giudizio. Se lista o utente mancano, il flag è non calcolabile.",
"data":"User ID dell’export, matrice ruoli/autorizzazioni, elenco account tecnici, periodo di validità degli accessi.",
"expected":"Modalità attiva/disattiva/non calcolabile, con motivazione. Se attivo, elenco degli utenti non autorizzati e delle relative scritture.",
"gap":"Non esiste un interruttore esplicito; lasciare la lista vuota non equivale metodologicamente a disattivare il test. Una whitelist generale non identifica utenti autorizzati ma inusuali per quel processo.",
"proposal":"Aggiungere stato del controllo: Attivo / Non applicabile / Dato non disponibile, con motivazione obbligatoria. Se attivo, peso forte 4 da sottoporre ad approvazione. Non trasformare l’assenza dei dati in esito negativo."},
{
"n":"9", "title":"Righe vuote o senza descrizione", "status":"PRESENTE PER LA DESCRIZIONE VUOTA", "score":"Peso iniziale attuale: 1; nel modello NORDSON valeva 4.",
"current":"La descrizione assente o composta solo da spazi attiva il flag. Il sistema non valuta righe completamente vuote, che normalmente vengono eliminate o causano problemi di importazione.",
"data":"Descrizione/causale di riga; idealmente anche testo di testata.",
"expected":"Flag “Descrizione vuota”; lista delle scritture prive di spiegazione.",
"gap":"Non intercetta descrizioni scarse ma presenti, quali “giroconto”, “rettifica”, “varie” o “adj”. Va chiarito se “riga vuota” significhi riga fisica dell’export oppure causale vuota.",
"proposal":"Mantenere descrizione vuota come rischio forte (peso 4 coerente con NORDSON, da approvare). Aggiungere separatamente “descrizione scarsa/generica”, con dizionario configurabile e revisione dei falsi positivi."},
{
"n":"10", "title":"Conto insolito o raro usato 2–4 volte all’anno", "status":"PRESENTE, MA LA SOGLIA VA DEFINITA", "score":"Peso opzionale attuale: non impostato; quindi il flag può esistere senza punti.",
"current":"Conta le righe per conto nella popolazione caricata e segnala frequenza < soglia. Con soglia 5 intercetta 1–4 utilizzi; non 0, perché un conto mai usato non compare nel giornale.",
"data":"Codice conto, popolazione dell’intero anno, piano dei conti e preferibilmente storico precedente.",
"expected":"Frequenza osservata, soglia applicata e flag conto raro.",
"gap":"“Da 2 a 4” escluderebbe i conti usati una sola volta, che sono potenzialmente ancora più rari. Il calcolo attuale conta righe, non scritture, e un caricamento trimestrale non rappresenta l’anno.",
"proposal":"Definire “da 1 a 4 utilizzi nell’intero esercizio” oppure motivare l’esclusione della frequenza 1. Contare scritture/documenti distinti oltre alle righe. Peso proposto da approvare: 4 per un segnale direttamente collegato all’ISA, con analisi storica quando disponibile."},
{
"n":"11", "title":"Test di sequenza numerica e controllo per pagina/serie", "status":"PRESENTE PER NUMERI PURAMENTE NUMERICI", "score":"Nessun punteggio: controllo separato.",
"current":"Ordina i numeri documento numerici, elimina duplicati e identifica i valori mancanti. Intervalli oltre 10.000 mancanti vengono riepilogati senza enumerazione.",
"data":"Numero protocollo reale, regole di numerazione, società, esercizio, sezionale/tipo documento. Per PDF bollato: numero pagina estratto e totale pagine.",
"expected":"Elenco gap per serie, più eventuali pagine mancanti/duplicate nel documento.",
"gap":"Le serie alfanumeriche sono escluse; serie diverse mescolate creano falsi gap. Il controllo pagina PDF non esiste. Un numero mancante può essere annullato legittimamente.",
"proposal":"Mantenere fuori dal punteggio. Configurare serie e reset; richiedere evidenza per numeri annullati. Aggiungere controllo pagine solo per documenti con numerazione attendibile e distinguere completezza del file da completezza contabile."},
{
"n":"12", "title":"Cifre finali ripetute", "status":"ASSENTE", "score":"Nessun peso attuale; da approvare.",
"current":"Non esiste un controllo dedicato. La divisibilità per 10 non intercetta correttamente finali come 77 o 99.",
"data":"Importo nella valuta originale e numero di decimali. Occorre definire se analizzare parte intera, centesimi o stringa normalizzata.",
"expected":"Flag e pattern rilevato, ad esempio ultime due cifre intere uguali (…77) o centesimi ripetuti (…,99). L’esempio 86.514,99 sarebbe rilevato per i centesimi 99.",
"gap":"La definizione “e così via” è troppo ampia: 77 può essere un importo intero o una terminazione; molti prezzi legittimi finiscono in ,99. Senza soglia minima si generano numerosi falsi positivi.",
"proposal":"Approvare pattern separati: ultime 2/3 cifre intere ripetute; centesimi 00/11/22…99; sequenze ripetute. Applicare una soglia minima e un peso basso 1, salvo combinazione con altri segnali."},
{
"n":"13", "title":"Seconda analisi limitata ai dati sospetti", "status":"ASSENTE COME FASE AUTONOMA", "score":"Nessun punteggio aggiuntivo.",
"current":"Il sistema consente di filtrare ed esportare le righe “da investigare”, ma non esegue un secondo motore analitico su tale sottoinsieme.",
"data":"Risultati della prima fase, scrittura completa, documenti di supporto, esito delle verifiche e motivazioni.",
"expected":"Coda di casi: raggruppamento per scrittura, ranking del rischio, confronto con scritture simili, richiesta evidenze, conclusione preparer/reviewer e tracciamento degli override.",
"gap":"Rieseguire le stesse regole sul sottoinsieme non aggiunge informazione e altera le distribuzioni (media/frequenze). Il sottoinsieme non deve diventare la nuova popolazione statistica.",
"proposal":"Realizzare un workflow di secondo livello, non un semplice rerun: mantenere le metriche della popolazione originale, arricchire le scritture selezionate e registrare esito finale (giustificata, errore, rettifica, escalation)."},
{
"n":"14", "title":"Numero di conto superiore a 10 caratteri/cifre", "status":"ASSENTE", "score":"Nessun peso attuale; da approvare.",
"current":"Il conto è conservato come testo ma la lunghezza non viene validata.",
"data":"Codice conto grezzo e specifica del piano dei conti/ERP. Va deciso se contare solo cifre o tutti i caratteri, escludendo spazi, separatori e prefissi legittimi.",
"expected":"Flag “formato conto anomalo”, lunghezza osservata e regola prevista per quella società.",
"gap":"Più di 10 caratteri non è universalmente anomalo: alcuni ERP usano segmenti, centri di costo concatenati o codici alfanumerici. Una regola globale produrrebbe falsi positivi.",
"proposal":"Trasformarlo in validazione configurabile del formato conto (regex/lunghezza/lookup nel piano dei conti), non in regola universale. Preferire “conto inesistente nel piano dei conti”. Peso basso 1 o controllo di qualità dati separato, da approvare."},
{
"n":"15", "title":"Keyword variabili, visura e soggetti con poteri", "status":"KEYWORD PRESENTI; LIBRERIA STANDARD E OCR VISURA ASSENTI", "score":"Keyword: peso iniziale attuale 1; nel modello NORDSON valevano 4. OCR: nessun peso.",
"current":"Il sistema cerca nella descrizione parole configurate manualmente, senza distinzione maiuscole/minuscole. Esiste anche una lista di conti infragruppo. Non carica né analizza la visura.",
"data":"Dizionario standard versionato, termini cliente, elenco parti correlate; visura aggiornata; nomi, codici fiscali, cariche e poteri; matrice degli accessi al gestionale e user ID.",
"expected":"Tre livelli separati: A) keyword di rischio nella causale; B) nominativi/cariche estratti dalla visura con conferma umana; C) incrocio tra persone, user ID, autorizzazioni e scritture. Ogni match deve mostrare origine e grado di affidabilità.",
"gap":"La visura dimostra cariche e poteri legali, non chi abbia accesso tecnico o chi possa materialmente manomettere il giornale. L’OCR può sbagliare nomi e ruoli. Parole generiche sulla frode producono molti falsi positivi; non serve necessariamente una libreria esterna.",
"proposal":"Creare un dizionario interno approvato e modificabile, con categorie (rettifica, override, urgenza, parti correlate, closing). OCR locale della visura con revisione obbligatoria campo per campo. Incrociare poi con accessi IT; mai etichettare automaticamente una persona come sospetta. Pesi distinti: keyword contestuale da approvare; match persona+utente non deve essere conclusivo senza evidenze."},
]

for c in controls:
    doc.add_heading(f"{c['n']}. {c['title']}", level=2)
    fill = GREEN if c["status"].startswith("PRESENTE") and "MA" not in c["status"] and "CON" not in c["status"] else AMBER if "PARZIALE" in c["status"] or "MA" in c["status"] or "CON" in c["status"] or "KEYWORD" in c["status"] else RED
    add_callout("Stato", c["status"] + ". " + c["score"], fill)
    add_label("Comportamento attuale", c["current"])
    add_label("Dati e fonte", c["data"])
    add_label("Risultato atteso", c["expected"])
    add_label("Gap e rischi", c["gap"])
    add_label("Comportamento TO-BE proposto", c["proposal"])

page_break()
doc.add_heading("3. Architettura funzionale proposta", level=1)
doc.add_paragraph("Per trasformare l’MVP in prodotto finito è utile separare le verifiche in livelli. In questo modo i risultati restano interpretabili e i controlli di completezza non vengono confusi con gli indicatori di frode.")
steps = [
    ("1. Qualità e completezza input", "Mappatura, formati, conti validi, quadratura Dare/Avere, sequenza, pagine e riconciliazione al bilancio di verifica."),
    ("2. Profilazione della popolazione", "Frequenze per conto/utente/tipo, valori medi, periodo coperto, valute, serie numeriche e completezza dei metadati."),
    ("3. Screening di rischio", "Applicazione dei controlli approvati a righe e scritture complete, mantenendo esiti vero/falso/non calcolabile."),
    ("4. Finestra di closing", "Selezione dedicata degli ultimi cinque giorni lavorativi e delle scritture create dopo la data di chiusura con competenza precedente."),
    ("5. Second-level review", "Raggruppamento dei sospetti, richiesta documenti, spiegazioni, autorizzazioni e conclusioni con doppia revisione."),
    ("6. Reporting e audit trail", "Parametri, fonti, versione delle regole, eccezioni, override, preparer/reviewer, timestamp ed export riproducibile."),
]
add_table(["Fase", "Contenuto"], steps, [2600, 7160], 8.8)

doc.add_heading("Regola di scoring raccomandata", level=2)
doc.add_paragraph("La seguente impostazione è una proposta per discussione, non una prescrizione normativa né una decisione già approvata:")
add_bullet("Indicatori deboli/contestuali: peso 1 (dimensione relativa, cifra tonda, weekend, fuori orario, cifre ripetute).")
add_bullet("Indicatori forti: peso 4 solo se sorretti da dati affidabili (retrodatazione significativa, utente non autorizzato, descrizione vuota, parte correlata/conto raro secondo metodologia).")
add_bullet("Controlli di completezza e qualità: nessun punteggio; esito separato bloccante o di eccezione.")
add_bullet("Scritture di closing: popolazione dedicata obbligatoria, indipendente dalla soglia aggregata.")
add_bullet("Non calcolabile: non assegna punti, ma deve generare un’avvertenza di copertura e comparire nel report.")

doc.add_heading("Dati minimi richiesti al cliente", level=2)
for x in [
    "Libro giornale completo con ID scrittura, numero documento, data effettiva, data/ora creazione, utente, conto, Dare/Avere, valuta e descrizione.",
    "Piano dei conti e bilanci di verifica iniziale/finale.",
    "Materialità e performance materiality approvate, con valuta e periodo.",
    "Calendario lavorativo, Paese/sede, data di chiusura e policy di closing.",
    "Matrice utenti/ruoli, account tecnici e storico delle autorizzazioni.",
    "Elenco parti correlate, conti infragruppo, visura aggiornata e struttura del gruppo.",
]: add_bullet(x)

doc.add_heading("Requisiti di governo", level=2)
for x in [
    "Versionare parametri, dizionari, calendari e regole; nessuna soglia nascosta.",
    "Richiedere motivazione per disattivare un controllo o accettare un dato mancante.",
    "Conservare l’origine di ogni riga e non alterare i documenti del cliente.",
    "Prevedere revisione umana dell’OCR e delle corrispondenze nominative.",
    "Proteggere visure, nomi, user ID e dati contabili con accessi, retention e log coerenti con la policy dello studio.",
]: add_bullet(x)

doc.add_heading("4. Priorità di sviluppo proposta", level=1)
priority = [
    ("P0 — metodologia", "Unità riga/scrittura; pesi; soglia; closing obbligatorio; definizioni di data, media e conto raro."),
    ("P1 — affidabilità", "Media automatica; calendario per Paese; data chiusura; toggle utenti; cifra tonda 10k/100k; esiti non calcolabili."),
    ("P2 — integrità", "Quadratura Dare/Avere, piano dei conti, serie numeriche/pagine, riconciliazione al bilancio di verifica."),
    ("P3 — approfondimento", "Cifre ripetute, dizionario keyword, workflow sospetti, dossier e audit trail."),
    ("P4 — document intelligence", "OCR visura, conferma umana, mapping persone–user ID–accessi e controlli privacy."),
]
add_table(["Priorità", "Contenuto"], priority, [2300, 7460], 8.8)

doc.add_heading("Criteri di accettazione del prodotto", level=2)
for x in [
    "Ogni risultato mostra valore osservato, soglia, formula, peso, fonte e motivazione.",
    "La stessa popolazione e configurazione producono lo stesso risultato ripetibile.",
    "I controlli non eseguiti sono conteggiati e visibili.",
    "I risultati sono disponibili sia per riga sia per scrittura completa.",
    "La finestra di closing è sempre analizzata e documentata.",
    "Il revisore può concludere e il reviewer può approvare o riaprire il caso.",
]: add_bullet(x)

page_break()
doc.add_heading("5. Conclusioni per l’approvazione", level=1)
doc.add_paragraph("Il nucleo dell’MVP è valido come strumento di screening, ma non è ancora sufficiente per qualificare il JET come prodotto completo e metodologicamente autosufficiente. Sette requisiti sono presenti in forma sostanziale, cinque sono parziali o diversi dalla richiesta e tre sono assenti come funzionalità autonome; inoltre alcune voci combinate richiedono una separazione concettuale.")
add_callout("Raccomandazione", "Approvare prima la metodologia e il dizionario dei controlli; solo dopo avviare lo sviluppo. La priorità non è aggiungere il maggior numero possibile di flag, ma garantire che ciascun risultato sia riproducibile, motivato e collegato a dati verificabili.", GREEN)
doc.add_heading("Delibere suggerite", level=2)
for x in [
    "Approvazione del perimetro dei 15 controlli, con separazione tra rischio, qualità dati e completezza.",
    "Approvazione dei pesi e dei controlli obbligatori indipendenti dal punteggio.",
    "Approvazione dei valori di default, lasciandoli modificabili e documentati per cliente.",
    "Nomina del responsabile metodologico e del responsabile della validazione tecnica.",
    "Avvio dello sviluppo secondo priorità P0–P4 e test su dati sintetici seguiti da validazione controllata su un caso reale.",
]: add_bullet(x)

doc.add_heading("6. Fonti e tracciabilità", level=1)
sources = [
    ("ISA Italia 240", "§33(a), §A38, §A43, §A44 e §A45, come riportati nel documento di studio fornito."),
    ("JET_libro_giornale_fondamento_normativo.pdf", "Documento di studio preliminare, 9 settembre 2026, 5 pagine."),
    ("00_riferimento_tecnico.md", "Sintesi della pipeline NORDSON/BTI, criteri, pesi e gap sul conto contabile."),
    ("Quadra — backend/jet/models.py", "Contratti dati e parametri attualmente disponibili."),
    ("Quadra — backend/jet/criteri.py", "Formule e calcolo effettivo dei flag/punteggi."),
    ("Quadra — backend/jet/sequenza.py", "Test attuale dei protocolli numerici mancanti."),
    ("Quadra — ui/src/jet/JetDashboard.tsx", "Caselle, pesi iniziali e soglia esposti oggi nell’interfaccia."),
]
add_table(["Fonte", "Utilizzo nella relazione"], sources, [3300, 6460], 8.4)
doc.add_paragraph("Nota: la relazione descrive lo stato osservato del progetto alla data indicata. I punteggi proposti sono raccomandazioni da sottoporre ad approvazione professionale; l’ISA Italia 240 non prescrive valori numerici specifici.")

# Page number field in footer
p = sec.footer.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("Pagina ")
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
r._r.addnext(fld)

# Keep table rows together where possible and repeat headers
for table in doc.tables:
    trPr = table.rows[0]._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); trPr.append(rep)
    for row in table.rows:
        pr = row._tr.get_or_add_trPr(); cant = OxmlElement("w:cantSplit"); pr.append(cant)

doc.core_properties.title = "Journal Entry Testing — Relazione AS-IS e TO-BE"
doc.core_properties.subject = "Valutazione dei controlli JET per evoluzione da MVP a prodotto"
doc.core_properties.author = "Quadra — relazione tecnica"
doc.core_properties.keywords = "JET, ISA 240, journal entry testing, audit, controlli"
doc.save(DOCX)
print(DOCX)
