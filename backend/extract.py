from __future__ import annotations

import csv
import io
import os
import re
from datetime import datetime
from pathlib import Path

TEXT_EXTS = {".txt", ".csv", ".tsv", ".xml", ".json", ".md"}
SHEET_EXTS = {".xlsx", ".xls", ".xlsm"}
PDF_EXTS = {".pdf"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}


def _safe_read_text_last_page(path: Path, limit: int = 20000) -> str:
    """Mastrini/giornale TXT: solo l'ultima pagina di stampa (o la coda del file)."""
    raw = path.read_bytes()
    text = ""
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        text = raw.decode("utf-8", errors="replace")
    if "\f" in text:
        parts = [p for p in text.split("\f") if p.strip()]
        return (parts[-1] if parts else text)[-limit:]
    marks = list(re.finditer(r"Pagina\s+\d+", text, re.I))
    if marks:
        start = text.rfind("\n", 0, marks[-1].start()) + 1
        return text[start:][-limit:]
    return text[-limit:]


def parse_mastrini_last(text: str) -> dict:
    """Dati di chiusura dall'ultima pagina mastrini/giornale."""
    page = None
    pm = list(re.finditer(r"Pagina\s+(\d+)", text, re.I))
    if pm:
        page = int(pm[-1].group(1))
    data_reg = None
    saldo_al = list(re.finditer(r"Progr\.e saldo al\s+(\d{1,2}/\d{1,2}/\d{2,4})", text, re.I))
    if saldo_al:
        data_reg = saldo_al[-1].group(1)
    else:
        dates = find_dates(text)
        if dates:
            data_reg = dates[-1]
    desc = ""
    sc = re.search(
        r"Sottoconto\s+\S+\s+(.+?)(?:\s{2,}Elaborazione|\s*$)",
        text,
        re.I | re.M,
    )
    if sc:
        desc = re.sub(r"\s+", " ", sc.group(1)).strip()
    cut = text[: saldo_al[-1].start()] if saldo_al else text
    dated = list(re.finditer(r"^\s*\d{1,2}/\d{1,2}/\d{2,4}\b", cut, re.M))
    if dated:
        bits = []
        for line in cut[dated[-1].start() :].splitlines():
            s = re.sub(r"\d{1,2}/\d{1,2}/\d{2,4}", " ", line)
            s = re.sub(r"[\d.]+,\d{2}", " ", s)
            s = re.sub(r"\b\d+\b", " ", s)
            s = re.sub(r"\s+", " ", s).strip(" -")
            if re.search(r"[A-Za-zÀ-ÿ]{4,}", s) and not re.search(
                r"Saldi preced|Progr\.e saldo|Pagina|Elaborazione", s, re.I
            ):
                bits.append(s)
        narrative = " ".join(bits)
        if narrative:
            desc = f"{desc} — {narrative}".strip(" —") if desc else narrative
    if not desc:
        for line in reversed(text.splitlines()):
            s = line.strip()
            if len(s) < 4 or s.startswith("-"):
                continue
            if re.search(r"FINE STAMPA|Saldi preced|Progr\.e saldo|Pagina\s+\d+|Elaborazione", s, re.I):
                continue
            if re.search(r"[A-Za-zÀ-ÿ]{4,}", s):
                desc = re.sub(r"\s+", " ", s)
                break
    return {"page": page, "data_reg": data_reg, "descrizione": desc[:80], "nr_reg": None}


def _safe_read_text(path: Path, limit: int = 80000) -> str:
    raw = path.read_bytes()[: limit * 2]
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            return raw.decode(enc)[:limit]
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")[:limit]


def extract_xlsx_text(path: Path, limit: int = 80000) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    chunks: list[str] = []
    try:
        # Prefer compact sheets (bilancino) over million-row giornali
        sheets = list(wb.worksheets)
        preferred = [s for s in sheets if re.search(r"verifica|bilancino|tabe|mapping", s.title, re.I)]
        order = preferred + [s for s in sheets if s not in preferred]
        for ws in order[:6]:
            chunks.append(f"# {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > 250:
                    break
                vals = [str(c) for c in row if c is not None]
                if vals:
                    chunks.append(" | ".join(vals))
                if sum(len(x) for x in chunks) > limit:
                    break
            if sum(len(x) for x in chunks) > limit:
                break
    finally:
        wb.close()
    return "\n".join(chunks)[:limit]


def extract_pdf_text(
    path: Path, limit: int = 80000, ocr: bool = True, layout: bool = False, pages: str = "key"
) -> tuple[str, str]:
    """Returns (text, method). PyMuPDF first (handles AES); OCR only if asked."""
    text = ""
    native_pages = 200 if pages == "all" else 40
    try:
        import pymupdf

        doc = pymupdf.open(str(path))
        try:
            if doc.needs_pass:
                return "", "pdf-encrypted"
            parts = []
            n = doc.page_count
            if pages == "last" and n:
                parts.append(doc[n - 1].get_text() or "")
            else:
                native_pages = 200 if pages == "all" else 40
                for i, page in enumerate(doc):
                    if i >= native_pages:
                        break
                    parts.append(page.get_text() or "")
            text = "\n".join(parts)
        finally:
            doc.close()
    except Exception:
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            parts = []
            if pages == "last" and reader.pages:
                parts.append(reader.pages[-1].extract_text() or "")
            else:
                cap = 200 if pages == "all" else native_pages
                for i, page in enumerate(reader.pages[:cap]):
                    parts.append(page.extract_text() or "")
            text = "\n".join(parts)
        except Exception:
            text = ""
    if len(text.strip()) >= 80:
        return text[:limit], "pdf-text"
    if ocr:
        ocr_limit = 200000 if pages == "all" else limit
        ocr_t = ocr_pdf(path, ocr_limit, layout=layout, pages=pages)
        if ocr_t.strip():
            return ocr_t, "ocr"
    return text[:limit], "pdf-text"


_OCR = None


def _rapidocr():
    global _OCR
    if _OCR is not None:
        return _OCR
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        try:
            from rapidocr import RapidOCR
        except ImportError:
            return None
    _OCR = RapidOCR()
    return _OCR


def ocr_wants_layout() -> bool:
    """VL 0.9B solo se richiesto. Su CPU Mac il default è PP-OCRv6 (leggero)."""
    return os.getenv("QUADRA_OCR_VL", "0").strip().lower() in {"1", "true", "yes"}


def _ocr_max_pages() -> int:
    try:
        return max(1, int(os.getenv("QUADRA_OCR_MAX_PAGES", "4")))
    except ValueError:
        return 4


def _ocr_deep_max() -> int:
    try:
        return max(_ocr_max_pages(), int(os.getenv("QUADRA_OCR_DEEP_MAX_PAGES", "80")))
    except ValueError:
        return 80


def _ocr_deep_skip_over() -> int:
    try:
        return max(_ocr_deep_max(), int(os.getenv("QUADRA_OCR_DEEP_SKIP_OVER", "80")))
    except ValueError:
        return 80


def _ocr_page_indices(n: int, max_pages: int) -> list[int]:
    if n <= max_pages:
        return list(range(n))
    head = max(1, max_pages // 2)
    tail = max_pages - head
    return sorted(set(list(range(head)) + list(range(n - tail, n))))


def ocr_image_bytes(data: bytes, layout: bool = False) -> str:
    layout = bool(layout and ocr_wants_layout())
    provider = os.getenv("QUADRA_OCR_PROVIDER", "paddle").lower()
    if provider in {"paddle", "auto"}:
        try:
            from .paddle_ocr import paddle_ocr

            text, _ = paddle_ocr(data, layout=layout)
            if text.strip():
                return text
        except Exception:
            # RapidOCR remains an offline-compatible fallback during migration.
            pass
    if provider == "none":
        return ""
    engine = _rapidocr()
    if engine is None:
        return ""
    result, _ = engine(data)
    if not result:
        return ""
    lines = []
    for row in result:
        if len(row) >= 2:
            lines.append(str(row[1]))
    return "\n".join(lines)


def ocr_pdf(
    path: Path,
    limit: int = 80000,
    max_pages: int | None = None,
    layout: bool = False,
    pages: str = "key",
) -> str:
    layout = bool(layout and ocr_wants_layout())
    try:
        import pymupdf
    except ImportError:
        return ""
    doc = pymupdf.open(str(path))
    texts = []
    try:
        n = doc.page_count
        if pages == "last":
            indices = [n - 1] if n else []
        elif pages == "all":
            cap = _ocr_deep_max()
            indices = list(range(min(n, cap)))
        else:
            if max_pages is None:
                max_pages = _ocr_max_pages()
            indices = _ocr_page_indices(n, max_pages)
        dpi = 160 if layout else 200
        for i in indices:
            page = doc[i]
            pix = page.get_pixmap(matrix=pymupdf.Matrix(dpi / 72, dpi / 72), alpha=False)
            texts.append(ocr_image_bytes(pix.tobytes("png"), layout=layout))
    finally:
        doc.close()
    return "\n".join(texts)[:limit]


def ocr_image_file(path: Path, layout: bool = False) -> str:
    return ocr_image_bytes(path.read_bytes(), layout=layout)[:80000]


def extract_file(
    path: Path, ocr: bool = True, layout: bool = False, pages: str = "key"
) -> tuple[str, str]:
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        if pages == "last":
            return _safe_read_text_last_page(path), "txt"
        return _safe_read_text(path), "txt"
    if ext in SHEET_EXTS:
        try:
            return extract_xlsx_text(path), "xlsx"
        except Exception as e:
            return str(e), "xlsx"
    if ext in PDF_EXTS:
        return extract_pdf_text(path, ocr=ocr, layout=layout, pages=pages)
    if ext in IMAGE_EXTS:
        if not ocr:
            return "", "ocr"
        text = ocr_image_file(path, layout=layout)
        return text, "ocr"
    return "", "filename"


def pdf_page_count(path: Path) -> int:
    try:
        import pymupdf

        doc = pymupdf.open(str(path))
        try:
            return int(doc.page_count)
        finally:
            doc.close()
    except Exception:
        return 0


def parse_amount(text: str) -> float | None:
    if text is None:
        return None
    s = str(text).strip()
    s = s.replace("€", "").replace("EUR", "").replace(" ", "")
    if re.match(r"^-?\d{1,3}(\.\d{3})+,\d{1,2}$", s):
        s = s.replace(".", "").replace(",", ".")
    elif "," in s and "." not in s:
        s = s.replace(",", ".")
    s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return float(s)
    except ValueError:
        return None


def find_amounts(text: str) -> list[float]:
    found = []
    for m in re.finditer(r"-?\d{1,3}(?:[.\s]\d{3})+,\d{2}|-?\d{1,3}(?:[.\s]\d{3})*,\d{2}|-?\d+[.,]\d{2}", text):
        val = parse_amount(m.group())
        if val is not None and abs(val) >= 0.01:
            found.append(val)
    return found


def find_f24_importo(text: str) -> float | None:
    """Total on an F24 quietanza. Prefer Italian thousands (629.911,47), not OCR scraps."""
    italian: list[float] = []
    for m in re.finditer(r"\d{1,3}(?:\.\d{3})+,\d{2}", text):
        val = parse_amount(m.group())
        if val is not None and val >= 1:
            italian.append(val)
    if italian:
        return italian[-1]
    labeled = []
    for line in text.splitlines():
        if not re.search(r"importo|versare|totale|erario|saldo", line, re.I):
            continue
        amts = find_amounts(line)
        if amts:
            labeled.append(max(amts, key=abs))
    if labeled:
        return max(labeled, key=abs)
    amts = [a for a in find_amounts(text) if a >= 10]
    return max(amts, key=abs) if amts else None


def find_statement_balance(text: str) -> float | None:
    """Closing balance from an e/c: prefer lines that say saldo, ignore tiny noise."""
    best = None
    for line in text.splitlines():
        low = line.lower()
        if not re.search(r"saldo|disponibilit|new balance|closing", low):
            continue
        if re.search(r"saldo iniziale|previous|precedente", low):
            continue
        amts = find_amounts(line)
        if not amts:
            continue
        val = amts[-1]
        if abs(val) < 1:
            continue
        best = val
    return best


def find_protocol(text: str) -> str | None:
    m = re.search(r"\b(\d{15,20})\b", text)
    return m.group(1) if m else None


def find_dates(text: str) -> list[str]:
    return re.findall(r"\b(\d{1,2}[./-]\d{1,2}[./-]\d{2,4})\b", text)


def find_payment_date(text: str) -> str | None:
    """Data versamento on an F24: compact OCR date near DATA DEL VERSAMENTO, else last date."""
    compact = re.sub(r"[\s.:]", "", text)
    upper = compact.upper()
    idx = upper.find("DATADELVERSAMENTO")
    windows = []
    if idx >= 0:
        windows.append(compact[max(0, idx - 48) : idx + 48])
    for chunk in windows:
        matches = re.findall(r"(0[1-9]|[12]\d|3[01])(0[1-9]|1[0-2])(20\d{2})", chunk)
        if matches:
            d, mo, y = matches[-1]
            return f"{d}/{mo}/{y}"
    dated_lines = []
    for line in text.splitlines():
        found = find_dates(line)
        if not found:
            continue
        if re.search(r"vers|delega|quietanza|pagat|protocol", line, re.I):
            return found[-1]
        dated_lines.extend(found)
    return dated_lines[-1] if dated_lines else None


def _as_int(value) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return None


def _is_dateish(value) -> bool:
    if isinstance(value, datetime):
        return True
    if isinstance(value, str) and re.match(r"\d{1,2}[./-]\d{1,2}[./-]\d{2,4}", value.strip()):
        return True
    return False


def parse_giornale_last(path: Path) -> dict:
    """Last registration on a SAP-style libro giornale xlsx (col A often empty)."""
    from openpyxl import load_workbook

    last_reg: list | None = None
    last_page = None
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            vals = [v for v in row if v is not None]
            blob = " ".join(str(v) for v in vals)
            pm = re.search(r"Pagina\s+(\d+)", blob, re.I)
            if pm:
                n = int(pm.group(1))
                last_page = n if last_page is None else max(last_page, n)
            if len(vals) < 2:
                continue
            nr = _as_int(vals[0])
            if nr is None or nr < 1:
                continue
            if _is_dateish(vals[1]):
                last_reg = vals
    finally:
        wb.close()
    if not last_reg:
        return {}
    # SAP: Numero | Data acq. | N.doc | Periodo | Data reg. | Data doc. | tipo | descrizione
    if len(last_reg) > 4 and _is_dateish(last_reg[4]):
        data_reg = last_reg[4]
    else:
        dates = [v for v in last_reg if _is_dateish(v)]
        data_reg = dates[-1] if dates else last_reg[1]
    desc = ""
    for i, cell in enumerate(last_reg):
        if isinstance(cell, str) and re.fullmatch(r"[A-Z]{2}", cell.strip()):
            for nxt in last_reg[i + 1 :]:
                if isinstance(nxt, str) and len(nxt.strip()) > 3 and not _is_dateish(nxt):
                    desc = nxt.strip()
                    break
            break
    if not desc:
        for cell in last_reg:
            if isinstance(cell, str) and len(cell) > 3 and not _is_dateish(cell):
                desc = cell.strip()
                break
    return {
        "nr_reg": _as_int(last_reg[0]),
        "data_reg": data_reg,
        "page": last_page if last_page and last_page > 1 else None,
        "descrizione": desc,
    }


def parse_csv_accounts(path: Path) -> list[dict]:
    text = _safe_read_text(path, 400000)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t,")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    rows = []
    for raw in reader:
        if not raw:
            continue
        keys = {k.lower().strip(): v for k, v in raw.items() if k}
        acc = (
            keys.get("conto")
            or keys.get("account")
            or keys.get("acc nb")
            or keys.get("codice")
            or keys.get("coge")
        )
        name = keys.get("descrizione") or keys.get("account name") or keys.get("nome") or ""
        amt = (
            keys.get("saldo")
            or keys.get("importo")
            or keys.get("amount")
            or keys.get("saldo periodo")
        )
        prior = keys.get("saldo n-1") or keys.get("prior") or keys.get("saldo precedente")
        if not acc:
            continue
        rows.append(
            {
                "account": str(acc).strip(),
                "name": str(name).strip(),
                "amount": parse_amount(amt) if amt else None,
                "prior": parse_amount(prior) if prior else None,
            }
        )
    return rows


def parse_xlsx_accounts(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = []
    try:
        ws = None
        for name in wb.sheetnames:
            if re.search(r"verifica|bilancino", name, re.I):
                ws = wb[name]
                break
        if ws is None:
            ws = wb.active
        header = None
        ita_idx = None
        acc_idx = 0
        name_idx = 1
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            vals = list(row)
            labels = [str(v).strip().lower() if v is not None else "" for v in vals]
            if header is None:
                if not any(h in labels for h in ("conto", "account", "acc nb")):
                    continue
                header = labels
                acc_idx = next(i for i, h in enumerate(header) if h in ("conto", "account", "acc nb", "codice"))
                name_idx = next(
                    (i for i, h in enumerate(header) if "descrizione" in h or "testo" in h or h in ("nome", "account name")),
                    acc_idx + 1,
                )
                if "ita" in header:
                    ita_idx = header.index("ita")
                continue
            acc = vals[acc_idx] if acc_idx < len(vals) else None
            if acc in (None, ""):
                continue
            name = vals[name_idx] if name_idx < len(vals) else ""
            amt_raw = vals[ita_idx] if ita_idx is not None and ita_idx < len(vals) else (vals[2] if len(vals) > 2 else None)
            amt = parse_amount(amt_raw) if amt_raw is not None else None
            if amt is None or amt == 0:
                continue
            rows.append(
                {
                    "account": str(acc).strip(),
                    "name": str(name or "").strip(),
                    "amount": amt,
                    "prior": None,
                }
            )
    finally:
        wb.close()
    return rows


def parse_docfinance_saldi(text: str) -> list[dict]:
    """Parse DocFinance 'Controllo situazione saldi' dumps."""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    out = []
    i = 0
    while i < len(lines) - 6:
        if re.match(r"^(BU|SA)\d{2}$", lines[i], re.I) and re.match(r"^[A-Z]{3}$", lines[i + 2]):
            banca = lines[i + 1]
            divisa = lines[i + 2]
            data = lines[i + 3]
            ric = parse_amount(lines[i + 4])
            calc = parse_amount(lines[i + 5])
            diff = parse_amount(lines[i + 6])
            if ric is not None or calc is not None:
                out.append(
                    {
                        "bank": banca,
                        "ccy": divisa,
                        "date": data,
                        "received": ric,
                        "booked": calc,
                        "diff": diff,
                    }
                )
            i += 7
            continue
        i += 1
    return out
