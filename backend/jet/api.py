"""API operativa del Journal Entry Testing."""

from __future__ import annotations

import hashlib
import re
from datetime import datetime, timezone
from pathlib import Path

from fastapi import APIRouter, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from backend.jet.criteri import (
    calcola_frequenza_conti,
    calcola_media_assoluta_registrazioni,
    valuta_riga,
)
from backend.jet.fonti import (
    ConfigurazioneDuplicatiJet,
    ConfigurazioneFonteJet,
    FonteJet,
    ImportazioneJet,
)
from backend.jet.ingest import leggi_righe_xlsx, mappa_righe_giornale
from backend.jet.ingest_pdf import estrai_righe_pdf, mappa_righe_pdf
from backend.jet.ingest_txt import (
    AnteprimaTxt,
    ispeziona_righe,
    ispeziona_txt,
    mappa_righe_txt,
    normalizza_intestazione,
)
from backend.jet.models import ParametriClienteJet, RigaGiornale
from backend.jet.pratica import (
    CreaPraticaJet,
    IntervalloSequenzaJet,
    MappaturaJet,
    PaginaRisultatiJet,
    PraticaJet,
    RisultatoJet,
)
from backend.jet.profilo import CreaProfiloEstrazione, ProfiloEstrazione
from backend.jet.sequenza import verifica_sequenza
from backend.jet.store import JetStore
from backend.workspace import output_dir, storage_root, write_inbox_file

router = APIRouter(prefix="/api/jet", tags=["jet"])
FORMATI_PROFILABILI = {".txt", ".pdf"}


def _store() -> JetStore:
    return JetStore(storage_root() / "jet.sqlite3")


def _pratica_or_404(pratica_id: str) -> PraticaJet:
    pratica = _store().get_pratica(pratica_id)
    if pratica is None:
        raise HTTPException(status_code=404, detail="Pratica JET non trovata")
    return pratica


def _file_path(pratica: PraticaJet) -> Path:
    fonte = _store().latest_fonte(pratica.id)
    if fonte is not None:
        return _fonte_path(fonte)
    if not pratica.file_originale_nome:
        raise HTTPException(
            status_code=409, detail="Carica prima il file Excel del libro giornale"
        )
    return (
        storage_root() / "pratiche" / pratica.id / "inbox" / pratica.file_originale_nome
    )


def _fonte_path(fonte: FonteJet) -> Path:
    return (
        storage_root()
        / "pratiche"
        / fonte.pratica_id
        / "inbox"
        / fonte.percorso_relativo
    )


def _fonte_or_404(pratica_id: str, fonte_id: str) -> FonteJet:
    fonte = _store().get_fonte(fonte_id)
    if fonte is None or fonte.pratica_id != pratica_id:
        raise HTTPException(status_code=404, detail="Fonte JET non trovata")
    return fonte


def _sync_pratica_fonti(
    pratica: PraticaJet, *, latest: FonteJet | None = None
) -> PraticaJet:
    fonti = _store().list_fonti(pratica.id)
    latest = latest or (fonti[-1] if fonti else None)
    update: dict[str, object] = {"numero_fonti": len(fonti)}
    if latest is not None:
        update.update(
            {
                "file_originale_nome": latest.nome_originale,
                "mappatura": latest.mappatura,
                "profilo_estrazione_id": latest.profilo_estrazione_id,
                "status": "file_caricato",
                "analizzato_at": None,
                "numero_registrazioni": 0,
                "numero_da_investigare": 0,
            }
        )
    elif pratica.file_originale_nome:
        update.update(
            {
                "file_originale_nome": None,
                "mappatura": None,
                "profilo_estrazione_id": None,
                "status": "parametri_configurati" if pratica.parametri else "bozza",
            }
        )
    pratica = pratica.model_copy(update=update)
    _store().save_pratica(pratica)
    return pratica


def _raw_rows(pratica: PraticaJet) -> list[dict]:
    suffix = Path(pratica.file_originale_nome or "").suffix.lower()
    if suffix in FORMATI_PROFILABILI:
        if not pratica.profilo_estrazione_id:
            raise HTTPException(
                status_code=409, detail="Applica prima un profilo di estrazione TXT"
            )
        profilo = _store().get_profilo(pratica.profilo_estrazione_id)
        if profilo is None:
            raise HTTPException(
                status_code=409, detail="Profilo di estrazione non trovato"
            )
        try:
            if suffix == ".pdf":
                return mappa_righe_pdf(_file_path(pratica), profilo)
            return mappa_righe_txt(_file_path(pratica), profilo)
        except (OSError, ValueError, TypeError) as exc:
            formato = "PDF" if suffix == ".pdf" else "TXT"
            raise HTTPException(
                status_code=400, detail=f"File {formato} non leggibile: {exc}"
            ) from exc
    try:
        return leggi_righe_xlsx(_file_path(pratica))
    except (OSError, ValueError, KeyError) as exc:
        raise HTTPException(
            status_code=400, detail=f"File Excel non leggibile: {exc}"
        ) from exc


def _anteprima_fonte(fonte: FonteJet) -> dict:
    path = _fonte_path(fonte)
    if fonte.formato == "xlsx":
        rows = leggi_righe_xlsx(path)
        headers = list(rows[0]) if rows else []
        if not headers:
            raise ValueError("Il file Excel non contiene intestazioni e righe dati")
        return {"intestazioni": headers}
    if fonte.formato == "pdf":
        anteprima = ispeziona_righe(estrai_righe_pdf(path), origine="pdf")
    else:
        anteprima = ispeziona_txt(path)
    return {
        "intestazioni": [],
        "intestazione": anteprima.intestazione,
        "riga_intestazione": anteprima.riga_intestazione,
        "righe_esempio": anteprima.righe_esempio,
        "codifica": anteprima.codifica,
        "profilo": _store().find_profilo_by_intestazione(
            normalizza_intestazione(anteprima.intestazione)
        ),
    }


def _righe_fonte(fonte: FonteJet) -> list[RigaGiornale]:
    path = _fonte_path(fonte)
    if fonte.formato == "xlsx":
        if fonte.mappatura is None:
            raise ValueError("Configura prima la mappatura Excel della fonte")
        return mappa_righe_giornale(leggi_righe_xlsx(path), fonte.mappatura)
    if fonte.profilo_estrazione_id is None:
        raise ValueError(
            f"Applica prima un profilo di estrazione {fonte.formato.upper()}"
        )
    profilo = _store().get_profilo(fonte.profilo_estrazione_id)
    if profilo is None:
        raise ValueError("Profilo di estrazione non trovato")
    if fonte.formato == "pdf":
        return mappa_righe_pdf(path, profilo)
    return mappa_righe_txt(path, profilo)


def _stage_fonte(fonte: FonteJet) -> FonteJet:
    try:
        fonte, _ = _store().replace_staging(fonte, _righe_fonte(fonte))
        return fonte
    except (OSError, ValueError, TypeError) as exc:
        errore = str(exc)
        _store().save_fonte(
            fonte.model_copy(update={"stato": "errore", "errore": errore})
        )
        raise HTTPException(
            status_code=400, detail=f"Fonte non importabile: {errore}"
        ) from exc


def _crea_fonte(
    pratica: PraticaJet, filename: str, data: bytes
) -> tuple[FonteJet, dict]:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".txt", ".pdf"}:
        raise HTTPException(
            status_code=400, detail="Sono accettati solo file .xlsx, .txt o .pdf"
        )
    fonte = FonteJet(
        pratica_id=pratica.id,
        nome_originale=filename,
        percorso_relativo="pending",
        formato=suffix[1:],
        sha256=hashlib.sha256(data).hexdigest(),
        dimensione_byte=len(data),
    )
    fonte = fonte.model_copy(update={"percorso_relativo": f"{fonte.id}/{filename}"})
    write_inbox_file(pratica.id, fonte.percorso_relativo, data)
    try:
        anteprima = _anteprima_fonte(fonte)
    except (OSError, ValueError) as exc:
        raise HTTPException(
            status_code=400,
            detail=f"File {fonte.formato.upper()} non leggibile: {exc}",
        ) from exc
    _store().save_fonte(fonte)
    pratica = _sync_pratica_fonti(pratica, latest=fonte)
    return fonte, {"pratica": pratica, "fonte": fonte, **anteprima}


@router.post("/pratiche", response_model=PraticaJet, status_code=201)
def create_pratica(body: CreaPraticaJet):
    pratica = PraticaJet(client=body.client.strip(), period=body.period.strip())
    _store().save_pratica(pratica)
    return pratica


@router.get("/pratiche", response_model=list[PraticaJet])
def list_pratiche():
    return _store().list_pratiche()


@router.get("/profili", response_model=list[ProfiloEstrazione])
def list_profili():
    return _store().list_profili()


@router.get("/profili/{profilo_id}", response_model=ProfiloEstrazione)
def get_profilo(profilo_id: str):
    profilo = _store().get_profilo(profilo_id)
    if profilo is None:
        raise HTTPException(status_code=404, detail="Profilo di estrazione non trovato")
    return profilo


@router.get("/pratiche/{pratica_id}", response_model=PraticaJet)
def get_pratica(pratica_id: str):
    return _pratica_or_404(pratica_id)


@router.put("/pratiche/{pratica_id}/parametri", response_model=PraticaJet)
def put_parametri(pratica_id: str, body: ParametriClienteJet):
    pratica = _pratica_or_404(pratica_id)
    status = "file_caricato" if pratica.file_originale_nome else "parametri_configurati"
    pratica = pratica.model_copy(
        update={
            "parametri": body,
            "status": status,
            "analizzato_at": None,
            "numero_registrazioni": 0,
            "numero_da_investigare": 0,
        }
    )
    _store().save_pratica(pratica)
    return pratica


@router.post("/pratiche/{pratica_id}/file")
async def upload_file(pratica_id: str, file: UploadFile = File(...)):
    pratica = _pratica_or_404(pratica_id)
    filename = Path(file.filename or "giornale.xlsx").name
    _, response = _crea_fonte(pratica, filename, await file.read())
    return response


@router.post("/pratiche/{pratica_id}/files")
async def upload_files(pratica_id: str, files: list[UploadFile] = File(...)):
    pratica = _pratica_or_404(pratica_id)
    caricate = []
    for file in files:
        filename = Path(file.filename or "giornale.xlsx").name
        fonte, response = _crea_fonte(pratica, filename, await file.read())
        pratica = response["pratica"]
        caricate.append(
            {
                "fonte": fonte,
                **{
                    key: value
                    for key, value in response.items()
                    if key not in {"pratica", "fonte"}
                },
            }
        )
    return {"pratica": pratica, "files": caricate}


@router.get("/pratiche/{pratica_id}/fonti", response_model=list[FonteJet])
def list_fonti(pratica_id: str):
    _pratica_or_404(pratica_id)
    return _store().list_fonti(pratica_id)


@router.get("/pratiche/{pratica_id}/importazioni", response_model=list[ImportazioneJet])
def list_importazioni(pratica_id: str):
    _pratica_or_404(pratica_id)
    return _store().list_importazioni(pratica_id)


@router.get("/pratiche/{pratica_id}/fonti/{fonte_id}/anteprima")
def preview_fonte(pratica_id: str, fonte_id: str):
    fonte = _fonte_or_404(pratica_id, fonte_id)
    try:
        return {"fonte": fonte, **_anteprima_fonte(fonte)}
    except (OSError, ValueError) as exc:
        raise HTTPException(
            status_code=400, detail=f"Fonte non leggibile: {exc}"
        ) from exc


@router.patch("/pratiche/{pratica_id}/fonti/{fonte_id}", response_model=FonteJet)
def configure_fonte(pratica_id: str, fonte_id: str, body: ConfigurazioneFonteJet):
    pratica = _pratica_or_404(pratica_id)
    fonte = _fonte_or_404(pratica_id, fonte_id)
    if body.attiva:
        configurata = (
            fonte.mappatura is not None or fonte.profilo_estrazione_id is not None
        )
        stato = "pronta" if fonte.numero_righe or configurata else "da_configurare"
    else:
        stato = "esclusa"
    fonte = fonte.model_copy(update={"attiva": body.attiva, "stato": stato})
    _store().save_fonte(fonte)
    _sync_pratica_fonti(pratica)
    return fonte


@router.delete("/pratiche/{pratica_id}/fonti/{fonte_id}", status_code=204)
def delete_fonte(pratica_id: str, fonte_id: str):
    pratica = _pratica_or_404(pratica_id)
    fonte = _fonte_or_404(pratica_id, fonte_id)
    removed = _store().delete_fonte(fonte.id)
    if removed is not None:
        _fonte_path(removed).unlink(missing_ok=True)
    _sync_pratica_fonti(pratica)


@router.put("/pratiche/{pratica_id}/fonti/{fonte_id}/file")
async def replace_fonte_file(
    pratica_id: str, fonte_id: str, file: UploadFile = File(...)
):
    pratica = _pratica_or_404(pratica_id)
    fonte = _fonte_or_404(pratica_id, fonte_id)
    filename = Path(file.filename or fonte.nome_originale).name
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".txt", ".pdf"}:
        raise HTTPException(
            status_code=400, detail="Sono accettati solo file .xlsx, .txt o .pdf"
        )
    data = await file.read()
    old_path = _fonte_path(fonte)
    replacement = fonte.model_copy(
        update={
            "nome_originale": filename,
            "percorso_relativo": f"{fonte.id}/{filename}",
            "formato": suffix[1:],
            "sha256": hashlib.sha256(data).hexdigest(),
            "dimensione_byte": len(data),
            "stato": "da_configurare",
            "mappatura": None,
            "profilo_estrazione_id": None,
            "numero_righe": 0,
            "errore": None,
        }
    )
    target = _fonte_path(replacement)
    candidate = target.with_name(f".{fonte.id}.candidate{suffix}")
    candidate.parent.mkdir(parents=True, exist_ok=True)
    candidate.write_bytes(data)
    try:
        preview_candidate = replacement.model_copy(
            update={
                "percorso_relativo": str(
                    candidate.relative_to(
                        storage_root() / "pratiche" / pratica.id / "inbox"
                    )
                )
            }
        )
        anteprima = _anteprima_fonte(preview_candidate)
        candidate.replace(target)
        if old_path != target:
            old_path.unlink(missing_ok=True)
    except (OSError, ValueError) as exc:
        candidate.unlink(missing_ok=True)
        raise HTTPException(
            status_code=400,
            detail=f"File {replacement.formato.upper()} non leggibile: {exc}",
        ) from exc
    _store().clear_staging(fonte.id)
    _store().save_fonte(replacement)
    pratica = _sync_pratica_fonti(pratica, latest=replacement)
    return {"pratica": pratica, "fonte": replacement, **anteprima}


@router.put("/pratiche/{pratica_id}/duplicati", response_model=PraticaJet)
def configure_duplicates(pratica_id: str, body: ConfigurazioneDuplicatiJet):
    pratica = _pratica_or_404(pratica_id)
    pratica = pratica.model_copy(
        update={
            "strategia_duplicati": body.strategia,
            "status": "file_caricato",
            "analizzato_at": None,
            "numero_registrazioni": 0,
            "numero_da_investigare": 0,
        }
    )
    _store().save_pratica(pratica)
    return pratica


@router.get("/pratiche/{pratica_id}/intestazioni")
def get_headers(pratica_id: str):
    pratica = _pratica_or_404(pratica_id)
    suffix = Path(pratica.file_originale_nome or "").suffix.lower()
    if suffix in FORMATI_PROFILABILI:
        try:
            if suffix == ".pdf":
                anteprima = ispeziona_righe(
                    estrai_righe_pdf(_file_path(pratica)), origine="pdf"
                )
            else:
                anteprima = ispeziona_txt(_file_path(pratica))
        except (OSError, ValueError) as exc:
            formato = "PDF" if suffix == ".pdf" else "TXT"
            raise HTTPException(
                status_code=400, detail=f"File {formato} non leggibile: {exc}"
            ) from exc
        return {
            "intestazioni": [],
            "intestazione": anteprima.intestazione,
            "riga_intestazione": anteprima.riga_intestazione,
            "righe_esempio": anteprima.righe_esempio,
            "codifica": anteprima.codifica,
            "profilo": _store().find_profilo_by_intestazione(
                normalizza_intestazione(anteprima.intestazione)
            ),
        }
    rows = _raw_rows(pratica)
    return {"intestazioni": list(rows[0]) if rows else []}


def _valida_posizioni_profilo(posizioni: dict[str, tuple[int, int]]) -> None:
    required = {"identificativo_registrazione", "data_effettiva"}
    if not required <= posizioni.keys():
        raise HTTPException(
            status_code=400,
            detail="Il profilo richiede identificativo_registrazione e data_effettiva",
        )
    if not ({"importo_netto", "importo_dare", "importo_avere"} & posizioni.keys()):
        raise HTTPException(
            status_code=400,
            detail="Mappare importo_netto oppure almeno una posizione Dare/Avere",
        )


def _valida_mappatura_excel(headers: set[str], mappatura: dict[str, str]) -> None:
    sconosciute = sorted(set(mappatura.values()) - headers)
    if sconosciute:
        raise HTTPException(
            status_code=400,
            detail=f"Colonne Excel non trovate: {', '.join(sconosciute)}",
        )
    campi_ammessi = set(RigaGiornale.model_fields) | {"importo_dare", "importo_avere"}
    campi_sconosciuti = sorted(set(mappatura) - campi_ammessi)
    if campi_sconosciuti:
        raise HTTPException(
            status_code=400,
            detail=f"Campi JET non riconosciuti: {', '.join(campi_sconosciuti)}",
        )
    required = {"identificativo_registrazione", "data_effettiva"}
    if not required <= mappatura.keys():
        raise HTTPException(
            status_code=400,
            detail="La mappatura richiede identificativo_registrazione e data_effettiva",
        )
    if not ({"importo_netto", "importo_dare", "importo_avere"} & mappatura.keys()):
        raise HTTPException(
            status_code=400,
            detail="Mappare importo_netto oppure almeno una colonna Dare/Avere",
        )


@router.put(
    "/pratiche/{pratica_id}/fonti/{fonte_id}/mappatura", response_model=FonteJet
)
def put_mappatura_fonte(pratica_id: str, fonte_id: str, body: MappaturaJet):
    pratica = _pratica_or_404(pratica_id)
    fonte = _fonte_or_404(pratica_id, fonte_id)
    if fonte.formato != "xlsx":
        raise HTTPException(status_code=400, detail="La fonte non è un file Excel")
    rows = leggi_righe_xlsx(_fonte_path(fonte))
    _valida_mappatura_excel(set(rows[0]) if rows else set(), body.mappatura)
    fonte = _stage_fonte(
        fonte.model_copy(
            update={
                "mappatura": body.mappatura,
                "profilo_estrazione_id": None,
            }
        )
    )
    _sync_pratica_fonti(pratica, latest=fonte)
    return fonte


def _anteprima_profilabile_fonte(fonte: FonteJet) -> AnteprimaTxt:
    if fonte.formato not in {"txt", "pdf"}:
        raise HTTPException(status_code=400, detail="La fonte non è un file TXT o PDF")
    preview = _anteprima_fonte(fonte)
    return AnteprimaTxt(
        codifica=preview["codifica"],
        riga_intestazione=preview["riga_intestazione"],
        intestazione=preview["intestazione"],
        righe_esempio=preview["righe_esempio"],
    )


@router.put(
    "/pratiche/{pratica_id}/fonti/{fonte_id}/profilo/{profilo_id}",
    response_model=FonteJet,
)
def apply_profilo_fonte(pratica_id: str, fonte_id: str, profilo_id: str):
    pratica = _pratica_or_404(pratica_id)
    fonte = _fonte_or_404(pratica_id, fonte_id)
    anteprima = _anteprima_profilabile_fonte(fonte)
    profilo = _store().get_profilo(profilo_id)
    if profilo is None:
        raise HTTPException(status_code=404, detail="Profilo di estrazione non trovato")
    if profilo.intestazione_riferimento != normalizza_intestazione(
        anteprima.intestazione
    ):
        raise HTTPException(
            status_code=400,
            detail="Il profilo non corrisponde all'intestazione della fonte",
        )
    _valida_posizioni_profilo(profilo.posizioni)
    fonte = fonte.model_copy(
        update={
            "profilo_estrazione_id": profilo.id,
            "mappatura": None,
            "stato": "pronta",
            "numero_righe": 0,
            "errore": None,
        }
    )
    _store().clear_staging(fonte.id)
    _store().save_fonte(fonte)
    _sync_pratica_fonti(pratica, latest=fonte)
    return fonte


@router.post(
    "/pratiche/{pratica_id}/fonti/{fonte_id}/profilo",
    response_model=FonteJet,
    status_code=201,
)
def create_and_apply_profilo_fonte(
    pratica_id: str, fonte_id: str, body: CreaProfiloEstrazione
):
    pratica = _pratica_or_404(pratica_id)
    fonte = _fonte_or_404(pratica_id, fonte_id)
    anteprima = _anteprima_profilabile_fonte(fonte)
    _valida_posizioni_profilo(body.posizioni)
    profilo = ProfiloEstrazione(
        nome=body.nome,
        riga_intestazione=anteprima.riga_intestazione,
        intestazione_riferimento=normalizza_intestazione(anteprima.intestazione),
        posizioni=body.posizioni,
    )
    _store().save_profilo(profilo)
    fonte = fonte.model_copy(
        update={
            "profilo_estrazione_id": profilo.id,
            "mappatura": None,
            "stato": "pronta",
            "numero_righe": 0,
            "errore": None,
        }
    )
    _store().clear_staging(fonte.id)
    _store().save_fonte(fonte)
    _sync_pratica_fonti(pratica, latest=fonte)
    return fonte


def _pratica_profilabile_or_400(pratica_id: str) -> tuple[PraticaJet, AnteprimaTxt]:
    pratica = _pratica_or_404(pratica_id)
    suffix = Path(pratica.file_originale_nome or "").suffix.lower()
    if suffix not in FORMATI_PROFILABILI:
        raise HTTPException(
            status_code=400, detail="La pratica non contiene un file TXT o PDF"
        )
    try:
        if suffix == ".pdf":
            anteprima = ispeziona_righe(
                estrai_righe_pdf(_file_path(pratica)), origine="pdf"
            )
        else:
            anteprima = ispeziona_txt(_file_path(pratica))
    except (OSError, ValueError) as exc:
        formato = "PDF" if suffix == ".pdf" else "TXT"
        raise HTTPException(
            status_code=400, detail=f"File {formato} non leggibile: {exc}"
        ) from exc
    return pratica, anteprima


@router.put("/pratiche/{pratica_id}/profilo/{profilo_id}", response_model=PraticaJet)
def apply_profilo(pratica_id: str, profilo_id: str):
    pratica, anteprima = _pratica_profilabile_or_400(pratica_id)
    profilo = _store().get_profilo(profilo_id)
    if profilo is None:
        raise HTTPException(status_code=404, detail="Profilo di estrazione non trovato")
    if profilo.intestazione_riferimento != normalizza_intestazione(
        anteprima.intestazione
    ):
        raise HTTPException(
            status_code=400,
            detail="Il profilo non corrisponde all'intestazione del file TXT",
        )
    _valida_posizioni_profilo(profilo.posizioni)
    fonte = _store().latest_fonte(pratica_id)
    if fonte is not None:
        fonte = fonte.model_copy(
            update={
                "profilo_estrazione_id": profilo.id,
                "mappatura": None,
                "stato": "pronta",
                "numero_righe": 0,
                "errore": None,
            }
        )
        _store().clear_staging(fonte.id)
        _store().save_fonte(fonte)
        return _sync_pratica_fonti(pratica, latest=fonte)
    pratica = pratica.model_copy(update={"profilo_estrazione_id": profilo.id})
    _store().save_pratica(pratica)
    return pratica


@router.post(
    "/pratiche/{pratica_id}/profilo", response_model=PraticaJet, status_code=201
)
def create_and_apply_profilo(pratica_id: str, body: CreaProfiloEstrazione):
    pratica, anteprima = _pratica_profilabile_or_400(pratica_id)
    _valida_posizioni_profilo(body.posizioni)
    profilo = ProfiloEstrazione(
        nome=body.nome.strip(),
        riga_intestazione=anteprima.riga_intestazione,
        intestazione_riferimento=normalizza_intestazione(anteprima.intestazione),
        posizioni=body.posizioni,
    )
    _store().save_profilo(profilo)
    fonte = _store().latest_fonte(pratica_id)
    if fonte is not None:
        fonte = fonte.model_copy(
            update={
                "profilo_estrazione_id": profilo.id,
                "mappatura": None,
                "stato": "pronta",
                "numero_righe": 0,
                "errore": None,
            }
        )
        _store().clear_staging(fonte.id)
        _store().save_fonte(fonte)
        return _sync_pratica_fonti(pratica, latest=fonte)
    pratica = pratica.model_copy(update={"profilo_estrazione_id": profilo.id})
    _store().save_pratica(pratica)
    return pratica


@router.put("/pratiche/{pratica_id}/mappatura", response_model=PraticaJet)
def put_mappatura(pratica_id: str, body: MappaturaJet):
    pratica = _pratica_or_404(pratica_id)
    rows = _raw_rows(pratica)
    headers = set(rows[0]) if rows else set()
    _valida_mappatura_excel(headers, body.mappatura)
    fonte = _store().latest_fonte(pratica_id)
    if fonte is not None:
        fonte = _stage_fonte(
            fonte.model_copy(
                update={
                    "mappatura": body.mappatura,
                    "profilo_estrazione_id": None,
                }
            )
        )
        return _sync_pratica_fonti(pratica, latest=fonte)
    pratica = pratica.model_copy(
        update={
            "mappatura": body.mappatura,
            "status": "file_caricato",
            "analizzato_at": None,
            "numero_registrazioni": 0,
            "numero_da_investigare": 0,
        }
    )
    _store().save_pratica(pratica)
    return pratica


@router.post("/pratiche/{pratica_id}/analizza", response_model=PraticaJet)
def analyze(pratica_id: str):
    pratica = _pratica_or_404(pratica_id)
    tutte_fonti = _store().list_fonti(pratica_id)
    fonti = [fonte for fonte in tutte_fonti if fonte.attiva]
    missing = []
    if pratica.parametri is None:
        missing.append("parametri")
    if tutte_fonti and not fonti:
        missing.append("almeno una fonte attiva")
    elif not fonti and pratica.file_originale_nome is None:
        missing.append("file Excel")
    if fonti:
        for fonte in fonti:
            if fonte.stato == "pronta":
                continue
            if fonte.stato == "errore":
                missing.append(f"fonte non valida: {fonte.nome_originale}")
            elif fonte.formato in {"txt", "pdf"}:
                missing.append(f"profilo di estrazione {fonte.formato.upper()}")
            else:
                missing.append("mappatura colonne")
    elif pratica.file_originale_nome:
        if Path(pratica.file_originale_nome).suffix.lower() in FORMATI_PROFILABILI:
            if pratica.profilo_estrazione_id is None:
                formato = Path(pratica.file_originale_nome).suffix[1:].upper()
                missing.append(f"profilo di estrazione {formato}")
        elif pratica.mappatura is None:
            missing.append("mappatura colonne")
    if missing:
        raise HTTPException(
            status_code=409,
            detail=f"Prima di analizzare completa: {', '.join(missing)}",
        )
    try:
        if fonti:
            for fonte in fonti:
                if fonte.numero_righe == 0:
                    _stage_fonte(fonte)
            staged_items = list(
                _store().iter_staged_rows(
                    pratica_id,
                    scarta_identiche=pratica.strategia_duplicati == "scarta_identiche",
                )
            )
            righe = [item.riga for item in staged_items]
        elif (
            Path(pratica.file_originale_nome or "").suffix.lower()
            in FORMATI_PROFILABILI
        ):
            righe = _raw_rows(pratica)
        else:
            righe = mappa_righe_giornale(_raw_rows(pratica), pratica.mappatura)
    except (ValueError, TypeError) as exc:
        raise HTTPException(
            status_code=400, detail=f"Errore nella mappatura dei dati: {exc}"
        ) from exc
    frequenze = calcola_frequenza_conti(righe)
    media_popolazione = calcola_media_assoluta_registrazioni(righe)
    esiti = [
        valuta_riga(r, pratica.parametri, frequenze, media_popolazione)
        for r in righe
    ]
    if fonti:
        risultati_persistenza = [
            RisultatoJet(
                riga=item.riga,
                esito=esito,
                fonte_id=item.fonte_id,
                numero_riga_fonte=item.numero_riga,
                hash_riga=item.hash_riga,
            )
            for item, esito in zip(staged_items, esiti, strict=True)
        ]
    else:
        risultati_persistenza = list(zip(righe, esiti, strict=True))
    mancanti, intervalli_raw = verifica_sequenza(righe)
    intervalli = [
        IntervalloSequenzaJet(precedente=a, successivo=b, quantita_mancanti=n)
        for a, b, n in intervalli_raw
    ]
    pratica = pratica.model_copy(
        update={
            "status": "analizzato",
            "analizzato_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "numero_registrazioni": len(esiti),
            "numero_da_investigare": sum(e.da_investigare for e in esiti),
            "valore_medio_registrazione_effettivo": (
                pratica.parametri.valore_medio_registrazione
                if pratica.parametri.valore_medio_registrazione is not None
                else media_popolazione
            ),
        }
    )
    _store().replace_analysis(pratica, risultati_persistenza, mancanti, intervalli)
    return pratica


def _ensure_analyzed(pratica_id: str) -> PraticaJet:
    pratica = _pratica_or_404(pratica_id)
    if pratica.status != "analizzato":
        raise HTTPException(
            status_code=409,
            detail="La pratica deve essere analizzata prima di consultare i risultati",
        )
    return pratica


@router.get("/pratiche/{pratica_id}/risultati", response_model=PaginaRisultatiJet)
def results(
    pratica_id: str,
    da_investigare: bool | None = None,
    conto_contabile: str | None = None,
    punteggio_minimo: int | None = Query(None, ge=0),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
):
    _ensure_analyzed(pratica_id)
    items, total = _store().query_results(
        pratica_id,
        da_investigare=da_investigare,
        conto_contabile=conto_contabile,
        punteggio_minimo=punteggio_minimo,
        limit=page_size,
        offset=(page - 1) * page_size,
    )
    return PaginaRisultatiJet(
        items=items,
        total=total,
        page=page,
        page_size=page_size,
        pages=(total + page_size - 1) // page_size,
    )


@router.get("/pratiche/{pratica_id}/sequenza")
def sequence(pratica_id: str):
    _ensure_analyzed(pratica_id)
    mancanti, intervalli = _store().sequence(pratica_id)
    return {"mancanti": mancanti, "intervalli_non_enumerati": intervalli}


FLAG_LABELS = {
    "flag_profit_impact": "Impatto sull'utile",
    "flag_oltre_dieci_volte_media": "> 10× media",
    "flag_sopra_performance_materiality": "Oltre performance materiality",
    "flag_importo_cifra_tonda": "Importo a cifra tonda",
    "flag_weekend": "Weekend",
    "flag_festivita": "Festività",
    "flag_fuori_orario": "Fuori orario",
    "flag_backdated": "Registrazione retrodatata",
    "flag_staff_non_autorizzato": "Staff non autorizzato",
    "flag_parte_correlata": "Parte correlata",
    "flag_descrizione_vuota": "Descrizione vuota",
    "flag_conto_insolito_raro": "Conto insolito/raro",
    "flag_conto_infragruppo_parte_correlata": "Conto infragruppo/parte correlata",
}


@router.get("/pratiche/{pratica_id}/export.xlsx")
def export_results(
    pratica_id: str,
    da_investigare: bool | None = None,
    conto_contabile: str | None = None,
    punteggio_minimo: int | None = Query(None, ge=0),
):
    pratica = _ensure_analyzed(pratica_id)
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Risultati JET"
    headers = [
        "ID registrazione",
        "Numero documento",
        "Data effettiva",
        "Conto contabile",
        "Importo netto",
        "Descrizione",
        "Utente",
        "Punteggio",
        "Da investigare",
        "Criteri scattati",
        "Fonte",
        "Riga fonte",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="31556F")
    nomi_fonti = {
        fonte.id: fonte.nome_originale for fonte in _store().list_fonti(pratica_id)
    }
    for item in _store().iter_export_results(
        pratica_id,
        da_investigare=da_investigare,
        conto_contabile=conto_contabile,
        punteggio_minimo=punteggio_minimo,
        batch_size=1000,
    ):
        r, e = item.riga, item.esito
        motivi = ", ".join(
            label for field, label in FLAG_LABELS.items() if getattr(e, field) is True
        )
        ws.append(
            [
                r.identificativo_registrazione,
                r.numero_documento,
                r.data_effettiva,
                r.conto_contabile,
                float(r.importo_netto),
                r.descrizione,
                r.utente,
                e.punteggio_totale,
                "Sì" if e.da_investigare else "No",
                motivi,
                nomi_fonti.get(item.fonte_id or "", ""),
                item.numero_riga_fonte,
            ]
        )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [20, 20, 16, 20, 18, 45, 22, 12, 16, 60, 28, 12]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    nome_sicuro = re.sub(
        r"[^\w.-]+", "_", f"JET_{pratica.client}_{pratica.period}"
    ).strip("._")
    path = output_dir(pratica_id) / f"{nome_sicuro or 'JET_risultati'}.xlsx"
    wb.save(path)
    return FileResponse(
        path,
        filename=path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
