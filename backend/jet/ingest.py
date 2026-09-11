"""Lettura tabellare e mappatura esplicita verso ``RigaGiornale``.

La lettura del file è separata dalla normalizzazione: il mapper accetta
semplici dizionari e non conosce né Excel né il cliente che li ha prodotti.
"""

from __future__ import annotations

import csv
from collections.abc import Iterable, Mapping
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from backend.jet.models import RigaGiornale

__all__ = [
    "data_o_none",
    "decimale_o_none",
    "leggi_righe_csv",
    "leggi_righe_tabellari",
    "leggi_righe_xlsx",
    "mappa_righe_giornale",
    "ora_o_none",
    "testo_o_none",
]


def _vuoto(valore: Any) -> bool:
    return valore is None or (isinstance(valore, str) and not valore.strip())


def testo_o_none(valore: Any) -> str | None:
    """Normalizza celle identificative senza trasformare un assente in ``''``."""
    if _vuoto(valore):
        return None
    if isinstance(valore, float) and valore.is_integer():
        return str(int(valore))
    return str(valore).strip()


def data_o_none(valore: Any) -> date | None:
    if _vuoto(valore):
        return None
    if isinstance(valore, datetime):
        return valore.date()
    if isinstance(valore, date):
        return valore
    testo = str(valore).strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(testo, formato).date()
        except ValueError:
            pass
    raise ValueError(f"Data non riconosciuta: {valore!r}")


def ora_o_none(valore: Any) -> time | None:
    if _vuoto(valore):
        return None
    if isinstance(valore, datetime):
        return valore.time()
    if isinstance(valore, time):
        return valore
    testo = str(valore).strip()
    for formato in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(testo, formato).time()
        except ValueError:
            pass
    raise ValueError(f"Ora non riconosciuta: {valore!r}")


def decimale_o_none(valore: Any) -> Decimal | None:
    if _vuoto(valore):
        return None
    if isinstance(valore, Decimal):
        return valore
    if isinstance(valore, float):
        return Decimal(str(valore))
    testo = str(valore).strip().replace(" ", "")
    if "," in testo:
        testo = testo.replace(".", "").replace(",", ".")
    try:
        return Decimal(testo)
    except InvalidOperation as exc:
        raise ValueError(f"Importo non riconosciuto: {valore!r}") from exc


def _valore(riga: Mapping[str, Any], mappatura: Mapping[str, str], campo: str) -> Any:
    colonna = mappatura.get(campo)
    return riga.get(colonna) if colonna is not None else None


def mappa_righe_giornale(
    righe: Iterable[Mapping[str, Any]],
    mappatura: Mapping[str, str],
) -> list[RigaGiornale]:
    """Mappa righe grezze sul contratto canonico usando solo scelte esplicite.

    La mappatura usa i nomi dei campi di ``RigaGiornale``. Per l'importo può
    contenere ``importo_dare`` e/o ``importo_avere``: se almeno uno dei due è
    configurato, il netto è sempre ricalcolato come Dare meno Avere e una
    cella vuota vale zero solo all'interno di questa sottrazione. In assenza
    di entrambi viene usata la colonna ``importo_netto`` già firmata.
    """
    risultato: list[RigaGiornale] = []
    usa_dare_avere = "importo_dare" in mappatura or "importo_avere" in mappatura

    for riga in righe:
        if usa_dare_avere:
            dare = decimale_o_none(_valore(riga, mappatura, "importo_dare"))
            avere = decimale_o_none(_valore(riga, mappatura, "importo_avere"))
            importo_netto = (dare or Decimal(0)) - (avere or Decimal(0))
        else:
            importo_netto = decimale_o_none(_valore(riga, mappatura, "importo_netto"))

        risultato.append(
            RigaGiornale(
                data_effettiva=data_o_none(_valore(riga, mappatura, "data_effettiva")),
                data_creazione=data_o_none(_valore(riga, mappatura, "data_creazione")),
                ora_creazione=ora_o_none(_valore(riga, mappatura, "ora_creazione")),
                identificativo_registrazione=testo_o_none(
                    _valore(riga, mappatura, "identificativo_registrazione")
                ),
                numero_documento=testo_o_none(
                    _valore(riga, mappatura, "numero_documento")
                ),
                importo_netto=importo_netto,
                descrizione=testo_o_none(_valore(riga, mappatura, "descrizione")),
                utente=testo_o_none(_valore(riga, mappatura, "utente")),
                conto_contabile=testo_o_none(
                    _valore(riga, mappatura, "conto_contabile")
                ),
            )
        )
    return risultato


def _intestazioni_univoche(valori: Iterable[Any]) -> list[str]:
    """Conserva intestazioni duplicate aggiungendo `` [2]``, `` [3]``..."""
    conteggi: dict[str, int] = {}
    risultato: list[str] = []
    for indice, valore in enumerate(valori, start=1):
        base = testo_o_none(valore) or f"Colonna {indice}"
        conteggi[base] = conteggi.get(base, 0) + 1
        occorrenza = conteggi[base]
        risultato.append(base if occorrenza == 1 else f"{base} [{occorrenza}]")
    return risultato


def leggi_righe_xlsx(
    percorso: str | Path,
    *,
    foglio: str | None = None,
    riga_intestazioni: int = 1,
) -> list[dict[str, Any]]:
    """Legge un foglio XLSX in dizionari, senza applicare mapping JET.

    Le intestazioni duplicate non possono convivere in un normale dizionario;
    vengono quindi rese indirizzabili in modo deterministico con suffissi
    `` [2]``, `` [3]``. Le righe completamente vuote sono ignorate.
    """
    workbook = load_workbook(Path(percorso), read_only=True, data_only=True)
    try:
        worksheet = workbook[foglio] if foglio is not None else workbook.active
        iteratore = worksheet.iter_rows(min_row=riga_intestazioni, values_only=True)
        intestazioni = _intestazioni_univoche(next(iteratore, ()))
        righe: list[dict[str, Any]] = []
        for valori in iteratore:
            if not any(not _vuoto(valore) for valore in valori):
                continue
            righe.append(dict(zip(intestazioni, valori)))
        return righe
    finally:
        workbook.close()


_ID_INTESTAZIONI = {
    "riga",
    "riga n",
    "riga n.",
    "transaction id",
    "transactionid",
    "id registrazione",
    "identificativo registrazione",
    "id",
}


def _assicura_colonna_riga(righe: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Se manca un identificativo, aggiunge ``Riga`` 1-based per la mappatura JET."""
    if not righe:
        return righe
    chiavi = {
        str(chiave).strip().lower().replace(".", "")
        for chiave in righe[0]
    }
    if chiavi & _ID_INTESTAZIONI:
        return righe
    return [{"Riga": str(indice), **riga} for indice, riga in enumerate(righe, start=1)]


def leggi_righe_csv(percorso: str | Path) -> list[dict[str, Any]]:
    """Legge un CSV in dizionari, senza applicare mapping JET.

    Accetta UTF-8 con o senza BOM. Se non c'è una colonna identificativo,
    aggiunge ``Riga`` (numero riga dati, da 1).
    """
    with Path(percorso).open(newline="", encoding="utf-8-sig") as handle:
        lettore = csv.reader(handle)
        grezze = list(lettore)
    if not grezze:
        return []
    intestazioni = _intestazioni_univoche(grezze[0])
    righe: list[dict[str, Any]] = []
    for valori in grezze[1:]:
        if not any(not _vuoto(valore) for valore in valori):
            continue
        padded = list(valori) + [""] * max(0, len(intestazioni) - len(valori))
        righe.append(dict(zip(intestazioni, padded[: len(intestazioni)])))
    return _assicura_colonna_riga(righe)


def leggi_righe_tabellari(percorso: str | Path) -> list[dict[str, Any]]:
    """Excel o CSV, stesso contratto di dizionari per la mappatura."""
    suffix = Path(percorso).suffix.lower()
    if suffix == ".csv":
        return leggi_righe_csv(percorso)
    return leggi_righe_xlsx(percorso)
