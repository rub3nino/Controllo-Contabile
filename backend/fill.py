from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill

from .catalog import DOCUMENT_NEED, ITEM_LABELS, MASTER_XLSX, SECTION_ITEMS, SECTION_TITLES, checklist_items
from .errors import UserError
from .extract import (
    extract_file,
    find_dates,
    find_f24_importo,
    find_payment_date,
    find_protocol,
    find_statement_balance,
    parse_csv_accounts,
    parse_docfinance_saldi,
    parse_giornale_last,
    parse_xlsx_accounts,
)
from .models import DocumentOut, PraticaIn, ProvenanceRow
from .provenance import ProvenanceLog

MONTHS_IT = {
    "gennaio": 1, "febbraio": 2, "marzo": 3, "aprile": 4, "maggio": 5, "giugno": 6,
    "luglio": 7, "agosto": 8, "settembre": 9, "ottobre": 10, "novembre": 11, "dicembre": 12,
    "gen": 1, "feb": 2, "mar": 3, "apr": 4, "mag": 5, "giu": 6,
    "lug": 7, "ago": 8, "set": 9, "ott": 10, "nov": 11, "dic": 12,
}

BANK_ALIASES = [
    ("intesa", "INTESA SANPAOLO SPA"),
    ("unicredit", "UNICREDIT SPA"),
    ("unicr", "UNICREDIT SPA"),
    ("unicred", "UNICREDIT SPA"),
    ("bpm", "BANCO BPM SPA"),
    ("bpn", "BANCO BPM / BPN"),
    ("bnl", "BNL GRUPPO BNP PARIBAS"),
    ("paribas", "BNP PARIBAS"),
    ("paschi", "MONTE DEI PASCHI DI SIENA"),
    ("mps", "MONTE DEI PASCHI DI SIENA"),
    ("piemonte", "BANCA DEL PIEMONTE S.P.A."),
    ("euromobiliare", "CREDEM EUROMOBILIARE SPA"),
    ("credem", "CREDITO EMILIANO SPA"),
    ("deutsche", "DEUTSCHE BANK"),
    ("mediobanca", "MEDIOBANCA SPA"),
    ("allianz", "ALLIANZ BANK SPA"),
    ("bper", "BPER BANCA SPA"),
    ("passadore", "BANCA PASSADORE"),
    ("sella", "BANCA SELLA SPA"),
    ("asti", "BANCA DI ASTI / BIVER"),
    ("biver", "BANCA DI ASTI / BIVER"),
    ("commerz", "COMMERZBANK"),
    ("bulbank", "UNICREDIT BULBANK"),
]


def parse_period_months(period: str) -> list[tuple[int, int]]:
    p = period.lower()
    year_m = re.search(r"(20\d{2})", p)
    year = int(year_m.group(1)) if year_m else datetime.now().year
    if re.search(r"\biii\b|\b3[°ºo]\b|\bterzo trim", p):
        return [(year, m) for m in (7, 8, 9)]
    if re.search(r"\bii\b|\b2[°ºo]\b|\bsecondo trim", p):
        return [(year, m) for m in (4, 5, 6)]
    if re.search(r"\biv\b|\b4[°ºo]\b|\bquarto trim", p):
        return [(year, m) for m in (10, 11, 12)]
    if re.search(r"\bi\b\s*trim|\b1[°ºo]\b|\bprimo trim", p):
        return [(year, m) for m in (1, 2, 3)]
    found = []
    for name, num in MONTHS_IT.items():
        if len(name) > 3 and name in p:
            found.append(num)
    found = sorted(set(found))
    if len(found) >= 2:
        start, end = found[0], found[-1]
        return [(year, m) for m in range(start, end + 1)]
    return [(year, datetime.now().month)]


def parse_date(value: str | None):
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(value[:10], fmt)
        except ValueError:
            continue
    return None


STATUS_FILL = {
    "✓": PatternFill("solid", fgColor="C6EFCE"),
    "wip": PatternFill("solid", fgColor="FFEB3B"),
    "N/A": PatternFill("solid", fgColor="D9D9D9"),
    "✗": PatternFill("solid", fgColor="FFC7CE"),
}

QUARTER_COLS = ("B", "C", "D", "E")  # 1°…4° trim on sheet C fondi/LIPE


def quarter_index(period: str) -> int:
    months = parse_period_months(period)
    m = months[0][1] if months else 4
    return min(4, max(1, (m - 1) // 3 + 1))


def quarter_year(period: str) -> int:
    months = parse_period_months(period)
    return months[0][0] if months else datetime.now().year


def adempimento_quarter(blob: str) -> int | None:
    """1..4 from filename like PREVIMODA_2TR26 / ENASARCO_1TR26. Longer roman first."""
    b = blob.lower()
    if re.search(r"(?:^|[_\s.-])4\s*tr|\biv\s*tr|4[°ºo]", b):
        return 4
    if re.search(r"(?:^|[_\s.-])3\s*tr|\biii\s*tr|3[°ºo]", b):
        return 3
    if re.search(r"(?:^|[_\s.-])2\s*tr|\bii\s*tr|2[°ºo]", b):
        return 2
    if re.search(r"(?:^|[_\s.-])1\s*tr|\bi\s*tr|1[°ºo]", b):
        return 1
    return None


def f24_month_from_name(name: str) -> int | None:
    m = re.search(
        r"(gen|feb|mar|apr|mag|giu|lug|ago|set|ott|nov|dic)[\s._-]?\d{2}",
        name.lower(),
    )
    if m:
        return MONTHS_IT.get(m.group(1))
    return None


def apply_status_fill(cell, status: str):
    fill = STATUS_FILL.get(status)
    if fill:
        cell.fill = fill


def _add_status_cf(ws, cell_range: str):
    """✓ green, wip/missing bright yellow, N/A gray, ✗ red."""
    for value, color in (
        ("✓", "C6EFCE"),
        ("wip", "FFEB3B"),
        ("N/A", "D9D9D9"),
        ("n/a", "D9D9D9"),
        ("✗", "FFC7CE"),
    ):
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        ws.conditional_formatting.add(
            cell_range,
            CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill),
        )


def _to_date(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        return parse_date(value.replace(".", "/"))
    return None


def set_cell(ws, coord: str, value, prov: ProvenanceLog, pratica: PraticaIn, **src):
    cell = ws[coord]
    if isinstance(cell, MergedCell):
        return None
    if value is None:
        return None
    cell.value = value
    return prov.add(
        client=pratica.client,
        period=pratica.period,
        sheet=ws.title,
        cell=coord,
        value=str(value),
        source_path=src.get("source_rel") or src.get("source_path", ""),
        source_rel=src.get("source_rel") or src.get("source_path", ""),
        source_name=src.get("source_name", "operatore"),
        page=src.get("page"),
        excerpt=src.get("excerpt", "")[:400],
        method=src.get("method", "umano"),
        confidence=src.get("confidence", 1.0),
        item_id=src.get("item_id"),
        human=src.get("human", False),
    )


class Filler:
    def __init__(self, pratica: PraticaIn, documents: list[DocumentOut], out_dir: Path, prov: ProvenanceLog):
        self.pratica = pratica
        self.documents = documents
        self.out_dir = out_dir
        self.prov = prov
        self.wb = None
        self.checklist: dict[str, str] = {}
        self.filled_mechanical = {"E": False, "C": False, "G": False, "B": False, "F": False}
        self.warnings: list[str] = []
        self.missing: list[dict] = []

    def docs_for(self, item_id: str) -> list[DocumentOut]:
        return [d for d in self.documents if d.item_id == item_id and not d.skip]

    def _src(self, doc: DocumentOut | None, item_id: str | None = None, extra: str = "") -> dict:
        if not doc:
            return {
                "source_path": "",
                "source_rel": "",
                "source_name": "regole/classificazione",
                "method": "filename",
                "excerpt": extra,
                "item_id": item_id,
                "confidence": 0.5,
            }
        return {
            "source_path": doc.rel or doc.path,
            "source_rel": doc.rel or doc.path,
            "source_name": doc.name,
            "method": doc.method,
            "excerpt": extra or doc.excerpt,
            "item_id": item_id or doc.item_id,
            "confidence": doc.confidence,
        }

    def fill_all(self, skip_items: list[str], na_items: list[str], skip_sections: list[str] | None = None) -> Path:
        dest = None
        for kind, payload in self.fill_iter(skip_items, na_items, skip_sections):
            if kind == "done":
                dest = payload
        assert dest is not None
        return dest

    def fill_iter(self, skip_items: list[str], na_items: list[str], skip_sections: list[str] | None = None):
        out_dir = self.out_dir
        out_dir.mkdir(parents=True, exist_ok=True)
        dest = out_dir / f"WPS_{_slug(self.pratica.client)}_{_slug(self.pratica.period)}.xlsx"
        if not MASTER_XLSX.exists():
            raise UserError(
                "Manca il modello Excel",
                "Non trovo Template_MASTER.xlsx nella cartella dell'applicazione. Senza quel file non si può compilare il WPS.",
                ["Template_MASTER.xlsx"],
            )
        shutil.copyfile(MASTER_XLSX, dest)
        self.wb = load_workbook(dest)
        skip = set(skip_items)
        na = set(na_items)
        skipped = set(skip_sections or self.pratica.skip_sections or [])

        jobs: list[tuple[str, str, object]] = [
            ("INDICE", "Apro il modello Excel e scrivo cliente e periodo", lambda: self._fill_indice_header()),
            ("Richiesta", "Compilo la checklist: segno cosa c'è e cosa manca", lambda: self._fill_richiesta(skip, na)),
        ]
        if "E" not in skipped:
            jobs.append(
                ("E", "Leggo estratti conto e riconciliazioni (i PDF scansionati richiedono OCR)", self._fill_banks)
            )
        if "C" not in skipped:
            jobs.append(("C", "Leggo quietanze F24 e adempimenti (OCR se il PDF non ha testo)", self._fill_f24))
        if "G" not in skipped:
            jobs.append(("G", "Compilo il bilancino di verifica, se presente", self._fill_trial_balance))
        if "B" not in skipped:
            jobs.append(("B", "Compilo i libri obbligatori (giornale, registri)", self._fill_books))
        if "F" not in skipped:
            jobs.append(("F", "Compilo i verbali degli organi sociali", self._fill_verbali))

        def _close():
            self._fill_judgment_notes(skipped)
            self._fill_indice_status(skip, na, skipped)
            self.prov.write_xlsx_sheet(self.wb)
            self.wb.save(dest)
            self._write_mancanti(out_dir / "mancanti.md")

        jobs.append(("INDICE", "Chiudo i semafori delle sezioni e salvo l'Excel", _close))

        total = len(jobs)
        for i, (sec, label, fn) in enumerate(jobs, start=1):
            yield "step", {"section": sec, "label": label, "step": i, "total": total}
            fn()
        yield "done", dest

    def _fill_indice_header(self):
        ws = self.wb["INDICE"]
        p = self.pratica
        manual = {"method": "umano", "source_name": "form pratica", "human": True, "confidence": 1}
        set_cell(ws, "A1", p.client, self.prov, p, **manual, item_id=None, excerpt="cliente")
        set_cell(ws, "M1", p.period, self.prov, p, **manual, excerpt="periodo")
        rd = parse_date(p.request_date)
        ad = parse_date(p.activity_date)
        if rd:
            c = ws["F5"]
            c.value = rd
            c.number_format = "dd/mm/yyyy"
            self.prov.add(
                client=p.client, period=p.period, sheet="INDICE", cell="F5",
                value=p.request_date, source_path="", source_name="form pratica",
                method="umano", excerpt="data invio richiesta", confidence=1, human=True,
            )
        if ad:
            c = ws["F6"]
            c.value = ad
            c.number_format = "dd/mm/yyyy"
            self.prov.add(
                client=p.client, period=p.period, sheet="INDICE", cell="F6",
                value=p.activity_date, source_path="", source_name="form pratica",
                method="umano", excerpt="data svolgimento", confidence=1, human=True,
            )
        if p.done_by:
            set_cell(ws, "C28", p.done_by, self.prov, p, **manual, excerpt="attività svolta da")
        if p.reviewed_by:
            set_cell(ws, "C29", p.reviewed_by, self.prov, p, **manual, excerpt="rivista da")
        rd_ws = self.wb["Richiesta doc"]
        set_cell(rd_ws, "N3", p.period, self.prov, p, **manual, excerpt="intestazione trimestre corrente")

    def _fill_richiesta(self, skip: set[str], na: set[str]):
        ws = self.wb["Richiesta doc"]
        for item in checklist_items():
            iid = item["id"]
            row = item["row"]
            docs = self.docs_for(iid)
            if iid == "C.3" and not docs:
                docs = [d for d in self.docs_for("C.1") if "collegio" in d.name.lower()]
            if iid in skip:
                status = "✗"
                note = "controllo saltato (operatore)"
                src = self._src(None, iid, note)
                src["human"] = True
                src["method"] = "umano"
            elif iid in na:
                status = "N/A"
                note = "non applicabile"
                src = self._src(None, iid, note)
                src["human"] = True
                src["method"] = "umano"
            elif docs:
                status = "✓"
                note = ", ".join(d.name for d in docs[:4])
                src = self._src(docs[0], iid, note)
            else:
                status = "wip"
                note = "documento non trovato in cartella"
                src = self._src(None, iid, note)
            self.checklist[iid] = status
            months = parse_period_months(self.pratica.period)
            if months:
                y, m = months[-1]
                label = f"{m:02d}/{y}"
                set_cell(ws, f"N{row}", label, self.prov, self.pratica, **src)
            set_cell(ws, f"P{row}", note[:80], self.prov, self.pratica, **src)
            qcell = ws[f"Q{row}"]
            qcell.value = status
            apply_status_fill(qcell, status)
            self.prov.add(
                client=self.pratica.client, period=self.pratica.period, sheet="Richiesta doc",
                cell=f"Q{row}", value=status, source_path=src.get("source_path", ""),
                source_name=src.get("source_name", ""), method=src.get("method", "filename"),
                excerpt=note[:200], confidence=src.get("confidence", 0.5), item_id=iid,
                human=src.get("human", False),
            )
        _add_status_cf(ws, "Q5:Q42")

    def _bank_docs_for_fill(self) -> list[DocumentOut]:
        """One statement per bank/account, preferring quarter-end files."""
        months = parse_period_months(self.pratica.period)
        last_m = months[-1][1] if months else 6
        month_names = {4: ("apr", "aprile"), 5: ("mag", "maggio", "may"), 6: ("giu", "giugno", "jun"),
                       3: ("mar", "marzo"), 9: ("set", "settembre"), 12: ("dic", "dicembre")}
        end_toks = month_names.get(last_m, ())
        ranked = []
        for doc in self.docs_for("F.1"):
            n = doc.name.lower()
            score = 0
            if any(t in n for t in end_toks) or re.search(rf"30[\s._-]?0?{last_m}|{last_m:02d}[\s._-]30|{last_m:02d}[\s._-]31", n):
                score += 12
            if re.search(r"01[\s._-]jul|01[\s._-]lug|luglio|july", n):
                score -= 20
            if any(x in n for x in ("aprile", "maggio", "apr-", "may-", "apr26", "mag26")) and last_m == 6:
                score -= 8
            if "scalar" in n:
                score -= 4
            if "ecommerce" in n or "e commerce" in n or "e-commerce" in n:
                score += 1
            ranked.append((score, doc))
        ranked.sort(key=lambda x: -x[0])
        seen = set()
        picked = []
        for score, doc in ranked:
            if score < 0:
                continue
            key = doc.name.lower()
            for alias, _ in BANK_ALIASES:
                if alias in key:
                    extra = "usd" if "usd" in key else ("ecom" if "ecomm" in key or "e comm" in key else "cc")
                    key = f"{alias}:{extra}"
                    break
            if key in seen:
                continue
            seen.add(key)
            picked.append(doc)
            if len(picked) >= 16:
                break
        return picked

    def _fill_banks(self):
        ws = self.wb["E"]
        p = self.pratica
        title = f"RICONCILIAZIONI BANCARIE — {p.period}"
        set_cell(ws, "A9", title, self.prov, p, source_name="periodo pratica", method="umano", excerpt=title, human=True)

        df_rows = []
        df_doc = None
        for doc in self.docs_for("F.2"):
            text, method = extract_file(Path(doc.path), ocr=True, layout=True)
            parsed = parse_docfinance_saldi(text)
            if parsed:
                df_rows = parsed
                df_doc = doc
                break

        docs = self._bank_docs_for_fill()
        row = 12
        used = 0
        for doc in docs:
            if row > 28:
                break
            text, method = extract_file(Path(doc.path), ocr=True, layout=True)
            blob_name = doc.name.lower()
            blob = f"{doc.name} {text}".lower()
            banca = "CONTO BANCARIO"
            alias_key = ""
            for alias, label in BANK_ALIASES:
                if alias in blob_name:
                    banca = label
                    alias_key = alias
                    break
            if not alias_key:
                for alias, label in BANK_ALIASES:
                    if alias in blob:
                        banca = label
                        alias_key = alias
                        break
            saldo = find_statement_balance(text)
            iban = re.search(r"IT\d{2}[A-Z]\d{10,}", text.replace(" ", ""), re.I)
            conto_m = re.search(r"conto\s*corrente\s*n\.?\s*([\d /]+)", text, re.I)
            conto = ""
            if conto_m:
                conto = re.sub(r"\s+", "", conto_m.group(1))[:20]
            elif iban:
                conto = iban.group(0)[-12:]
            ccy = "USD" if "usd" in blob else "cc"
            src = self._src(doc, "F.1", f"saldo={saldo}")
            src["method"] = method
            set_cell(ws, f"B{row}", banca, self.prov, p, **src)
            set_cell(ws, f"C{row}", ccy, self.prov, p, **src)
            if conto:
                set_cell(ws, f"D{row}", conto, self.prov, p, **src)
            if saldo is not None:
                ws[f"F{row}"].value = saldo
                ws[f"F{row}"].number_format = '#,##0.00'
                self.prov.add(
                    client=p.client, period=p.period, sheet="E", cell=f"F{row}",
                    value=str(saldo), source_rel=doc.rel, source_name=doc.name,
                    method=method, excerpt=doc.excerpt, confidence=doc.confidence, item_id="F.1",
                )
            if df_rows and alias_key:
                for rec in df_rows:
                    if alias_key in rec["bank"].lower() and rec.get("booked") is not None:
                        if ccy == "USD" and rec.get("ccy") != "USD":
                            continue
                        if ccy != "USD" and rec.get("ccy") == "USD":
                            continue
                        ws[f"E{row}"].value = rec["booked"]
                        ws[f"E{row}"].number_format = '#,##0.00'
                        self.prov.add(
                            client=p.client, period=p.period, sheet="E", cell=f"E{row}",
                            value=str(rec["booked"]), source_path=df_doc.rel if df_doc else "",
                            source_rel=df_doc.rel if df_doc else "",
                            source_name=df_doc.name if df_doc else "DocFinance",
                            method="pdf-text", excerpt=f"{rec['bank']} {rec.get('date')}", confidence=0.85, item_id="F.2",
                        )
                        break
            row += 1
            used += 1
        if used:
            self.filled_mechanical["E"] = True

        tb_docs = self.docs_for("B.1")
        accounts = []
        for doc in tb_docs:
            accounts.extend(self._load_accounts(doc))
        if accounts:
            for r in range(12, 29):
                if ws[f"E{r}"].value not in (None, ""):
                    continue
                banca = ws[f"B{r}"].value
                if not banca:
                    continue
                key = str(banca).split()[0].lower()
                for acc in accounts:
                    name = (acc.get("name") or "").lower()
                    if key and key in name and acc.get("amount") is not None:
                        ws[f"E{r}"].value = acc["amount"]
                        ws[f"E{r}"].number_format = '#,##0.00'
                        src_doc = tb_docs[0]
                        self.prov.add(
                            client=p.client, period=p.period, sheet="E", cell=f"E{r}",
                            value=str(acc["amount"]), source_path=src_doc.rel,
                            source_rel=src_doc.rel, source_name=src_doc.name, method="xlsx",
                            excerpt=f"{acc['account']} {acc['name']}", confidence=0.7, item_id="B.1",
                        )
                        break

        for r in range(12, 29):
            banca = ws[f"B{r}"].value
            e_val, f_val = ws[f"E{r}"].value, ws[f"F{r}"].value
            if banca and isinstance(e_val, (int, float)) and isinstance(f_val, (int, float)):
                delta = float(e_val) - float(f_val)
                if abs(delta) >= 0.005:
                    msg = f"{banca}: check bancario ≠ 0 (Δ {delta:.2f})"
                    self.warnings.append(msg)
                    self.prov.add(
                        client=p.client, period=p.period, sheet="E", cell=f"H{r}",
                        value=f"{delta:.2f}", source_path="",
                        source_name="saldo e/c vs co.ge",
                        method="formula", excerpt=msg, confidence=1.0, item_id="F.1",
                    )

    def _load_accounts(self, doc: DocumentOut) -> list[dict]:
        path = Path(doc.path)
        ext = path.suffix.lower()
        try:
            if ext in {".csv", ".txt", ".tsv"}:
                return parse_csv_accounts(path)
            if ext in {".xlsx", ".xlsm", ".xls"}:
                return parse_xlsx_accounts(path)
        except Exception:
            return []
        return []

    def _fill_f24(self):
        ws = self.wb["C"]
        p = self.pratica
        months = parse_period_months(p.period)
        cols = ["B", "C", "D", "E", "F", "G"]
        for i, (year, month) in enumerate(months[:6]):
            col = cols[i]
            dt = datetime(year, month, 1)
            cell = ws[f"{col}10"]
            cell.value = dt
            cell.number_format = "mmm-yy"
            self.prov.add(
                client=p.client, period=p.period, sheet="C", cell=f"{col}10",
                value=dt.strftime("%Y-%m"), source_path="", source_name="periodo pratica",
                method="umano", excerpt="mese F24 del trimestre", confidence=1, human=True, item_id="E.1",
            )
        all_f24 = self.docs_for("E.1")
        iva_docs = [d for d in all_f24 if re.search(r"iva", d.name, re.I)]
        docs = iva_docs or all_f24
        if len(all_f24) > len(docs):
            self.warnings.append(f"{len(all_f24)} quietanze F24 in cartella; in tabella C usate {len(docs)} IVA di periodo")
        used = 0
        for doc in docs:
            text, method = extract_file(Path(doc.path), ocr=False)
            if len(text.strip()) < 40:
                text, method = extract_file(Path(doc.path), ocr=True, layout=True)
            blob = f"{doc.name} {text}".lower()
            proto = find_protocol(text)
            dates = find_dates(text)
            name_month = f24_month_from_name(doc.name)
            target_col = None
            for i, (year, month) in enumerate(months[:6]):
                if name_month == month:
                    target_col = cols[i]
                    break
                long_names = [k for k, v in MONTHS_IT.items() if v == month and len(k) > 3]
                if any(n in doc.name.lower() for n in long_names):
                    target_col = cols[i]
                    break
                if any(n in blob for n in long_names):
                    target_col = cols[i]
                    break
                for d in dates:
                    parsed = parse_date(d.replace(".", "/"))
                    if parsed and parsed.month == month:
                        target_col = cols[i]
                        break
                if target_col:
                    break
            if target_col is None and used < len(months):
                target_col = cols[used]
            if not target_col:
                continue
            src = self._src(doc, "E.1")
            src["method"] = method
            pay = find_payment_date(text)
            dv = parse_date(pay.replace(".", "/")) if pay else None
            if dv:
                ws[f"{target_col}11"].value = dv
                ws[f"{target_col}11"].number_format = "dd/mm/yyyy"
                self.prov.add(
                    client=p.client, period=p.period, sheet="C", cell=f"{target_col}11",
                    value=pay, source_rel=doc.rel, source_name=doc.name,
                    method=method, excerpt=text[:200], confidence=doc.confidence, item_id="E.1",
                )
            importo = find_f24_importo(text)
            if importo is not None:
                ws[f"{target_col}12"].value = importo
                ws[f"{target_col}12"].number_format = '#,##0.00'
                self.prov.add(
                    client=p.client, period=p.period, sheet="C", cell=f"{target_col}12",
                    value=str(importo), source_rel=doc.rel, source_name=doc.name,
                    method=method, excerpt=text[:200], confidence=doc.confidence, item_id="E.1",
                )
            if proto:
                set_cell(ws, f"{target_col}13", proto, self.prov, p, **src)
            used += 1
        if used:
            self.filled_mechanical["C"] = True

        qcol = QUARTER_COLS[quarter_index(p.period) - 1]
        year = quarter_year(p.period)
        for i, col in enumerate(QUARTER_COLS, start=1):
            ws[f"{col}18"].value = f"{i}° trim. {year}"
        if self.docs_for("E.3"):
            set_cell(ws, f"{qcol}24", "✓", self.prov, p, **self._src(self.docs_for("E.3")[0], "E.3"))
        for doc in self.docs_for("E.2"):
            blob = (doc.name + " " + doc.excerpt).lower()
            dest_col = qcol
            qn = adempimento_quarter(blob)
            if qn:
                dest_col = QUARTER_COLS[qn - 1]
            for r in range(19, 24):
                label = str(ws[f"A{r}"].value or "").lower()
                if label in ("", "[fondo]"):
                    name = Path(doc.name).stem
                    if name.lower().endswith("pdf"):
                        name = name[:-3]
                    set_cell(ws, f"A{r}", name[:40], self.prov, p, **self._src(doc, "E.2"))
                    set_cell(ws, f"{dest_col}{r}", "✓", self.prov, p, **self._src(doc, "E.2"))
                    break
                if any(tok in blob for tok in label.split() if len(tok) > 3):
                    set_cell(ws, f"{dest_col}{r}", "✓", self.prov, p, **self._src(doc, "E.2"))
                    break

    def _fill_trial_balance(self):
        ws = self.wb["G"]
        p = self.pratica
        docs = self.docs_for("B.1")
        if not docs:
            return
        doc = docs[0]
        accounts = self._load_accounts(doc)
        if not accounts:
            self.warnings.append("B.1 presente ma senza saldi utilizzabili (colonna ITA vuota o file di mapping)")
            set_cell(
                ws, "A9",
                f"File {doc.name}: nessun saldo ITA da incollare. Tabella G non compilata.",
                self.prov, p, **self._src(doc, "B.1"),
            )
            return
        r = 10
        for acc in accounts[:480]:
            src = self._src(doc, "B.1", f"{acc['account']} {acc.get('amount')}")
            set_cell(ws, f"B{r}", acc["account"], self.prov, p, **src)
            set_cell(ws, f"D{r}", acc.get("name") or "", self.prov, p, **src)
            if acc.get("amount") is not None:
                ws[f"E{r}"].value = acc["amount"]
                ws[f"E{r}"].number_format = '#,##0.00'
                self.prov.add(
                    client=p.client, period=p.period, sheet="G", cell=f"E{r}",
                    value=str(acc["amount"]), source_rel=doc.rel, source_name=doc.name,
                    method=doc.method, excerpt=src["excerpt"], confidence=doc.confidence, item_id="B.1",
                )
            if acc.get("prior") is not None:
                ws[f"F{r}"].value = acc["prior"]
                ws[f"F{r}"].number_format = '#,##0.00'
            if acc.get("amount") is not None and acc.get("prior") is not None:
                ws[f"G{r}"] = f"=E{r}-F{r}"
                ws[f"H{r}"] = f'=IF(F{r}=0,"",E{r}/F{r}-1)'
                ws[f"H{r}"].number_format = "0.0%"
            r += 1
        if accounts:
            self.filled_mechanical["G"] = True
            set_cell(
                ws, "E9", p.period, self.prov, p,
                source_name="periodo pratica", method="umano", excerpt="header saldo periodo", human=True, item_id="B.1",
            )

    def _fill_books(self):
        ws = self.wb["B"]
        p = self.pratica
        ws["H10"].value = "Nome file"
        ws["H10"].font = Font(bold=True)
        ws.column_dimensions["H"].width = 42
        mapping = [
            (11, "D.1"),
            (12, "B.4"),
            (14, "D.3"),
            (15, "G.1"),
        ]
        for row, iid in mapping:
            docs = self.docs_for(iid)
            if not docs:
                continue
            doc = docs[0]
            src = self._src(doc, iid)
            set_cell(ws, f"H{row}", doc.name, self.prov, p, **src)
            wrap_b = Alignment(wrap_text=True, vertical="center")
            ws[f"H{row}"].alignment = wrap_b
            ws.row_dimensions[row].height = 32
            meta = {}
            path = Path(doc.path)
            if iid in ("B.4", "D.1") and path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
                try:
                    meta = parse_giornale_last(path)
                except Exception:
                    meta = {}
            text, method = extract_file(path, ocr=True)
            src["method"] = method
            if meta.get("nr_reg") is not None:
                set_cell(ws, f"C{row}", meta["nr_reg"], self.prov, p, **src)
            if meta.get("data_reg") is not None:
                dv = _to_date(str(meta["data_reg"]) if not isinstance(meta["data_reg"], datetime) else meta["data_reg"])
                if dv:
                    ws[f"D{row}"].value = dv
                    ws[f"D{row}"].number_format = "dd/mm/yyyy"
                    self.prov.add(
                        client=p.client, period=p.period, sheet="B", cell=f"D{row}",
                        value=str(meta["data_reg"]), source_rel=doc.rel, source_name=doc.name,
                        method=method, excerpt="data ultima registrazione", confidence=doc.confidence, item_id=iid,
                    )
            if meta.get("page") is not None:
                set_cell(ws, f"E{row}", meta["page"], self.prov, p, **src)
            if meta.get("descrizione"):
                set_cell(ws, f"F{row}", str(meta["descrizione"])[:80], self.prov, p, **src)
                ws[f"F{row}"].alignment = wrap_b
            elif not ws[f"F{row}"].value:
                dates = find_dates(text)
                if dates:
                    dv = parse_date(dates[-1].replace(".", "/"))
                    if dv and not ws[f"D{row}"].value:
                        ws[f"D{row}"].value = dv
                        ws[f"D{row}"].number_format = "dd/mm/yyyy"
                pag = re.search(r"pag(?:ina|\.)?\s*(\d+)", f"{doc.name} {text}", re.I)
                if pag and not ws[f"E{row}"].value:
                    set_cell(ws, f"E{row}", int(pag.group(1)), self.prov, p, **src)
            if not ws[f"G{row}"].value:
                set_cell(ws, f"G{row}", "ultimo aggiornamento rilevato dal file", self.prov, p, **src)
            self.filled_mechanical["B"] = True

    def _fill_verbali(self):
        ws = self.wb["F"]
        p = self.pratica
        ws.column_dimensions["B"].width = 48
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 22
        ws.column_dimensions["E"].width = 24
        ws.column_dimensions["F"].width = 52
        wrap = Alignment(wrap_text=True, vertical="center")
        for r in range(9, 14):
            ws.row_dimensions[r].height = 36
            for col in "BCDEF":
                ws[f"{col}{r}"].alignment = wrap
        mapping = [(11, "C.1"), (13, "C.2"), (10, "C.4")]
        for row, iid in mapping:
            docs = self.docs_for(iid)
            if not docs:
                continue
            doc = docs[0]
            text, method = extract_file(Path(doc.path), ocr=True)
            dates = find_dates(text)
            src = self._src(doc, iid)
            src["method"] = method
            cell_f = ws[f"F{row}"]
            cell_f.value = doc.name
            cell_f.alignment = wrap
            ws.row_dimensions[row].height = 36
            self.prov.add(
                client=p.client, period=p.period, sheet="F", cell=f"F{row}",
                value=doc.name, source_rel=doc.rel, source_name=doc.name,
                method=method, excerpt=doc.excerpt, confidence=doc.confidence, item_id=iid,
            )
            date_s = None
            mdate = re.search(r"(20\d{2})(\d{2})(\d{2})", doc.name)
            if mdate:
                date_s = f"{mdate.group(3)}/{mdate.group(2)}/{mdate.group(1)}"
            elif dates:
                date_s = dates[-1]
            dv = parse_date(date_s.replace(".", "/")) if date_s else None
            if dv:
                ws[f"D{row}"].value = dv
                ws[f"D{row}"].number_format = "dd/mm/yyyy"
                ws[f"D{row}"].alignment = wrap
                self.prov.add(
                    client=p.client, period=p.period, sheet="F", cell=f"D{row}",
                    value=date_s, source_rel=doc.rel, source_name=doc.name,
                    method=method, excerpt=text[:160], confidence=doc.confidence, item_id=iid,
                )
            pag = re.search(r"pag(?:ina|\.)?\s*(\d+)", f"{doc.name} {text}", re.I)
            if pag:
                set_cell(ws, f"C{row}", int(pag.group(1)), self.prov, p, **src)
                ws[f"C{row}"].alignment = wrap
            self.filled_mechanical["F"] = True
        docs_col = self.docs_for("C.3")
        if not docs_col:
            docs_col = [d for d in self.docs_for("C.1") if "collegio" in d.name.lower()]
        if docs_col:
            doc = docs_col[0]
            note = f"Collegio: {doc.name}"
            existing = ws["F11"].value
            ws["F11"].value = f"{existing} | {note}" if existing and "Collegio" not in str(existing) else (existing or note)
            ws["F11"].alignment = wrap
            ws.row_dimensions[11].height = 48
            self.prov.add(
                client=p.client, period=p.period, sheet="F", cell="F11",
                value=str(ws["F11"].value), source_rel=doc.rel, source_name=doc.name,
                method=doc.method, excerpt="nomina/verbale collegio sullo stesso PDF assemblea",
                confidence=doc.confidence, item_id="C.3",
            )

    def _fill_judgment_notes(self, skipped: set[str] | None = None):
        """A/D/H/I: elenca evidenze, senza conclusioni inventate."""
        skipped = skipped or set()
        mapping = {
            "A": ["A.1", "A.2", "A.3", "A.4"],
            "D": ["B.4"],
            "H": [],
            "I": ["A.2", "A.4"],
        }
        p = self.pratica
        for sheet, items in mapping.items():
            if sheet in skipped:
                continue
            found_docs: list[DocumentOut] = []
            labels: list[str] = []
            for iid in items:
                for doc in self.docs_for(iid):
                    found_docs.append(doc)
                    labels.append(f"{iid} {doc.name}")
            if labels:
                note = "Evidenza in cartella (nessuna conclusione): " + "; ".join(labels[:8])
                src_doc = found_docs[0]
            else:
                note = "Nessuna evidenza in cartella; chiusura e testo a mano."
                src_doc = None
            set_cell(
                self.wb[sheet], "A9", note, self.prov, p,
                **self._src(src_doc, items[0] if items else None, note),
            )

    def _fill_indice_status(self, skip: set[str], na: set[str], skip_sections: set[str] | None = None):
        ws = self.wb["INDICE"]
        p = self.pratica
        skipped = skip_sections or set()
        cells = {"A": "F10", "B": "F11", "C": "F12", "D": "F13", "E": "F14", "F": "F15", "G": "F16", "H": "F17", "I": "F18"}
        for sec, cell in cells.items():
            items = SECTION_ITEMS[sec]
            excerpt = ""
            if sec in skipped:
                status = "✗"
                excerpt = "sezione non selezionata dall'operatore"
            elif sec == "H" or not items:
                status = "wip"
                excerpt = "nessun documento in checklist; chiusura umana"
            else:
                statuses = []
                for iid in items:
                    if iid in skip:
                        statuses.append("✗")
                    elif iid in na:
                        statuses.append("N/A")
                    else:
                        statuses.append(self.checklist.get(iid, "wip"))
                excerpt = f"item {items} → {statuses}"
                if sec in ("A", "D", "I"):
                    status = "wip"
                    excerpt = f"sezione di giudizio; {excerpt}"
                elif all(s == "✓" or s in ("✗", "N/A") for s in statuses) and any(s == "✓" for s in statuses):
                    mech_ok = True
                    if sec in ("E", "C", "G") and not self.filled_mechanical.get(sec):
                        mech_ok = False
                    status = "✓" if mech_ok else "wip"
                else:
                    status = "wip"
            set_cell(
                ws, cell, status, self.prov, p,
                source_name="operatore" if sec in skipped else "regola INDICE",
                method="umano" if sec in skipped else "filename",
                excerpt=excerpt, item_id=None, confidence=1.0 if sec in skipped else 0.9,
                human=sec in skipped,
            )
            apply_status_fill(ws[cell], status)
        _add_status_cf(ws, "F10:F18")

    def _write_mancanti(self, path: Path):
        lines = [
            f"# Documenti mancanti",
            f"",
            f"Cliente: {self.pratica.client}",
            f"Periodo: {self.pratica.period}",
            f"",
        ]
        missing = []
        for item in checklist_items():
            st = self.checklist.get(item["id"], "wip")
            if st in ("wip", "", None):
                missing.append(item)
                need = DOCUMENT_NEED.get(item["id"], "Documento non trovato in cartella.")
                lines.append(f"- **{item['id']}** {ITEM_LABELS.get(item['id'], '')}")
                lines.append(f"  {need}")
        if not missing:
            lines.append("Nessun documento mancante (niente in wip).")
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        self.missing = [
            {
                "id": it["id"],
                "label": ITEM_LABELS.get(it["id"], it["id"]),
                "status": "wip",
                "need": DOCUMENT_NEED.get(it["id"], "Documento non trovato in cartella."),
            }
            for it in missing
        ]


def _slug(s: str) -> str:
    s = re.sub(r"[^\w]+", "_", s.strip(), flags=re.U)
    return s.strip("_")[:60] or "pratica"
